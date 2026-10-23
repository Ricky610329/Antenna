# -*- coding: utf-8 -*-
"""給日月光的「歷來實驗總整理」第二版投影片產生器，約 30 頁。

和初版 make_pptx.py 的差別：前三代改依碩論定稿重寫成自己的頁面、多放碩論原圖，數字只用
notes/genN_碩論對照.md 核對表判定一致的值；第四代改成四任務爬升圖加一頁量化總表；措辭照
docs/唯一真實.md §1.4、§1.5 的對外用語，條件、出處與限制一律寫進備忘稿。

版式沿用初版元件：16:9、白底、30pt 粗體標題、標題下灰細線加短深藍線、標題下一句描述、
小標 20pt 粗體深藍、表格淺藍表頭；不放底部結語帶、不放圖下小字。
第二輪依日月光視角審稿與數字稽核修訂；第四代的圖都取自 figs/gen4_ase/，原樣放上，
濾波器規格遮罩由 matplotlib 現畫。

跑法：在 repo 根執行 python docs/進度報告/10月_歷來實驗總整理/make_pptx_v2.py
→ 同夾 歷來實驗總整理_v2.pptx；本夾在 .gitignore，不進 git。
爬升圖由 script/figs/report_ase_climb.py 產生，時間軸圖由 script/figs/report_generations.py 產生。
"""
import math
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
FIG = HERE / "figs"
MEET = REPO / "docs" / "進度報告" / "組會_2026-10" / "figs"
OUT = HERE / "歷來實驗總整理_v2.pptx"

FOOTER_BAND = False     # 底部結語帶；第二版一律不放
SOURCE_LINE = False     # 頁底出處小字；第二版一律不放，出處寫進備忘稿

FONT = "Microsoft JhengHei"
INK = RGBColor(0x15, 0x2C, 0x3D)
BLUE = RGBColor(0x24, 0x5B, 0x85)
MUTED = RGBColor(0x52, 0x6A, 0x7D)
RULE = RGBColor(0xCC, 0xD8, 0xE2)
HEAD = RGBColor(0xEE, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x3A, 0x8A, 0x7C)
NAVY = RGBColor(0x24, 0x4C, 0x6A)
ORANGE = RGBColor(0xC0, 0x62, 0x2F)
GRAY = RGBColor(0xCB, 0xD3, 0xDA)
W, H = 13.333, 7.5
L, R = 0.583, 12.75
CW = R - L
BODY_BOTTOM = 6.88

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def repo(rel):
    return REPO / rel


def one(paths):
    """glob 結果必須恰好一個檔。檔名含研究紀錄代號的圖用 glob 取，原始碼裡不出現代號。"""
    paths = sorted(paths)
    assert len(paths) == 1, paths
    return paths[0]


# ---------------------------------------------------------------- 基本元件，沿用初版 make_pptx.py
def _run(par, text, size, bold=False, color=INK):
    r = par.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    rpr = r._r.get_or_add_rPr()                      # 中文字也指定正黑體，避免 PowerPoint 換成佈景字型
    rpr.set("lang", "zh-TW")                         # 繁中斷行規則：右括號不放行首
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)
    return r


def textbox(slide, x, y, w, h, lines, size=19, color=INK, align=None, anchor=None, spacing=1.15, after=0):
    """lines: list of str 或 (str, bold) 或 (str, bold, color)。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.line_spacing = spacing
        if after:
            par.space_after = Pt(after)
        if align is not None:
            par.alignment = align
        item = line if isinstance(line, tuple) else (line, False)
        _run(par, item[0], size, item[1], item[2] if len(item) > 2 else color)
    return tb


def rich_text(slide, x, y, w, h, blocks, size=16, head_size=20, anchor=MSO_ANCHOR.MIDDLE):
    """右欄：blocks＝[(藍色小標, [條列...]), ...]；小標 20pt 粗體深藍，條列前加「・」。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = anchor
    first = True
    for head, items in blocks:
        if head:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            par.line_spacing = 1.1
            par.space_after = Pt(4)
            if par is not tf.paragraphs[0]:
                par.space_before = Pt(14)
            _run(par, head, head_size, True, BLUE)
        for it in items:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            par.line_spacing = 1.12
            par.space_after = Pt(6)
            ppr = par._p.get_or_add_pPr()                # 懸掛縮排：換行後對齊「・」後的文字
            ppr.set("marL", str(int(Pt(size))))
            ppr.set("indent", str(-int(Pt(size))))
            _run(par, "・" + it, size, False, INK)
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    ln = s.line._get_or_add_ln()                     # 不畫外框
    ln.append(ln.makeelement(qn("a:noFill"), {}))
    s.shadow.inherit = False
    return s


def _as_stream(img):
    if isinstance(img, Image.Image):
        bio = BytesIO()
        img.save(bio, "PNG")
        bio.seek(0)
        return bio, img.size
    return str(img), Image.open(img).size


def trim(img, thresh=242, pad=10):
    """裁掉圖四周的白邊，讓圖在頁面上放得更大；只裁邊，不動內容。"""
    im = (img if isinstance(img, Image.Image) else Image.open(img)).convert("RGB")
    gray = im.convert("L").point(lambda v: 255 if v < thresh else 0)
    box = gray.getbbox()
    if not box:
        return im
    x0, y0, x1, y1 = box
    return im.crop((max(0, x0 - pad), max(0, y0 - pad), min(im.width, x1 + pad), min(im.height, y1 + pad)))


def picture_fit(slide, img, x, y, w, h, valign="center", trimmed=False):
    """把圖放進 (x,y,w,h) 框內、保持比例、水平置中；img 可以是路徑或 PIL Image。"""
    if trimmed:
        img = trim(img)
    src, (iw, ih) = _as_stream(img)
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    top = y if valign == "top" else y + (h - ph) / 2
    return slide.shapes.add_picture(src, Inches(x + (w - pw) / 2), Inches(top), Inches(pw), Inches(ph))


def page_number(slide):
    n = len(prs.slides)
    textbox(slide, 12.15, 7.06, 0.8, 0.32, [str(n)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def set_notes(slide, notes):
    if not notes:
        return
    text = "\n".join(notes) if isinstance(notes, (list, tuple)) else notes
    slide.notes_slide.notes_text_frame.text = text


def _text_width(s, pt):
    """粗估字串寬度，單位吋：全形字 1 em、半形約 0.55 em。"""
    return sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in s) * pt / 72


def content_slide(title, desc=None, notes=None):
    """回傳 (slide, 內容區頂端 y)。描述超過一行時內容區自動下移。"""
    s = prs.slides.add_slide(BLANK)
    textbox(s, L - 0.04, 0.30, CW, 0.70, [(title, True)], size=30, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, L, 1.07, CW, 0.012, RULE)
    rect(s, L, 1.055, 1.36, 0.04, BLUE)
    top = 1.75
    if desc:
        nline = math.ceil(_text_width(desc, 17) / (CW - 0.2))
        textbox(s, L - 0.04, 1.16, CW, 0.40 * nline + 0.1, [desc], size=17, color=INK)
        top = 1.80 if nline <= 1 else (2.12 if nline == 2 else 2.45)
    page_number(s)
    set_notes(s, notes)
    return s, top


def _border(cell, color=None, width_pt=0.75):
    """只畫下框線；左右上清空。框線元素要排在填色元素之前，這是 DrawingML 規定的順序。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(qn(tag), {"w": str(int(width_pt * 12700)), "cap": "flat", "cmpd": "sng",
                                        "algn": "ctr"})
        if tag == "a:lnB" and color is not None:
            fill = ln.makeelement(qn("a:solidFill"), {})
            clr = fill.makeelement(qn("a:srgbClr"), {"val": str(color)})
            fill.append(clr)
            ln.append(fill)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.insert(i, ln)


def table(slide, x, y, rows, col_w, header_size=15, body_size=14, row_h=0.5, bold_first_col=True, header_h=None,
          center_cols=()):
    """rows[0] 是表頭；儲存格給 str 或 (str, bold)；給 None 表示與左邊儲存格合併。"""
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_r))
    t = shp.table
    t.horz_banding = False
    t.first_row = True
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches((header_h or min(row_h, 0.5)) if i == 0 else row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            _border(cell, RULE if i < n_r - 1 else BLUE, 0.75 if i else 1.0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD if i == 0 else WHITE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if val is None:
                continue
            tf = cell.text_frame
            tf.word_wrap = True
            text = val[0] if isinstance(val, tuple) else val
            bold = i == 0 or (isinstance(val, tuple) and val[1]) or (bold_first_col and j == 0)
            for k, part in enumerate(text.split("\n")):
                par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                if k == 0:
                    par.text = ""
                par.line_spacing = 1.0
                if j in center_cols and not (j + 1 < n_c and row[j + 1] is None):
                    par.alignment = PP_ALIGN.CENTER
                _run(par, part, header_size if i == 0 else body_size, bold, INK)
        # 合併：連續的 None 併進左邊第一個有值的儲存格
        j = 0
        while j < n_c:
            if row[j] is not None:
                k = j
                while k + 1 < n_c and row[k + 1] is None:
                    k += 1
                if k > j:
                    t.cell(i, j).merge(t.cell(i, k))
                j = k + 1
            else:
                j += 1
    return shp


# ---------------------------------------------------------------- 頁型
def table_slide(title, desc, rows, col_w, notes=None, **kw):
    s, top = content_slide(title, desc, notes)
    table(s, L, top + 0.08, rows, col_w, **kw)
    return s


def fig_slide(title, desc, img, notes=None, pad=0.0, trimmed=False):
    s, top = content_slide(title, desc, notes)
    picture_fit(s, img, L + pad, top, CW - 2 * pad, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    return s


def fig_text_slide(title, desc, img, blocks, notes=None, img_w=6.6, size=16, trimmed=True):
    """左圖右文；blocks＝[(藍色小標, [條列...]), ...]。圖預設先裁白邊。"""
    s, top = content_slide(title, desc, notes)
    if isinstance(img, Image.Image):
        trimmed = False
    picture_fit(s, img, L, top, img_w, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    tx = L + img_w + 0.35
    rich_text(s, tx, top, R - tx, BODY_BOTTOM - top, blocks, size=size)
    return s


def columns_slide(title, desc, cols, notes=None, size=17):
    """cols: [(小標, [條列...]), ...]，兩到三欄文字。"""
    s, top = content_slide(title, desc, notes)
    n = len(cols)
    gap = 0.45
    cw = (CW - gap * (n - 1)) / n
    for k, (head, items) in enumerate(cols):
        x = L + k * (cw + gap)
        rect(s, x, top + 0.25, cw, 0.05, BLUE)
        rich_text(s, x, top + 0.4, cw, BODY_BOTTOM - top - 0.6, [(head, items)], size=size,
                  anchor=MSO_ANCHOR.TOP)
    return s


def fig_table_slide(title, desc, img, rows, col_w, notes=None, img_w=5.0, trimmed=True, **kw):
    """左圖右表：用在合併頁，圖與表各自精簡。"""
    s, top = content_slide(title, desc, notes)
    picture_fit(s, img, L, top, img_w, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    shp = table(s, L + img_w + 0.3, top + 0.1, rows, col_w, **kw)
    return s, shp


# ---------------------------------------------------------------- 濾波器規格遮罩
CJK = "C:/Windows/Fonts/msjh.ttc"


def spec_filter():
    """濾波器規格遮罩：只畫 S21 的規格位置，淺色是不允許的範圍；不放示範曲線。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font_manager.fontManager.addfont(CJK)
    plt.rcParams.update({"font.family": [font_manager.FontProperties(fname=CJK).get_name(), "DejaVu Sans"], "font.size": 15,
                         "axes.edgecolor": "#8A9AA8", "axes.labelcolor": "#152C3D", "xtick.color": "#152C3D",
                         "ytick.color": "#152C3D"})
    shade, ink = "#F3DCD3", "#152C3D"
    fig, axes = plt.subplots(1, 2, figsize=(10.4, 5.0), dpi=220)
    panels = [("雙埠濾波器 S21 (dB)", (24, 32), (25.5, 30.5), [(24, 25), (31, 32)], [24, 26, 28, 30, 32]),
              ("寬頻濾波器 S21 (dB)", (11, 40), (26, 30), [(11, 19), (36, 40)], [11, 15, 20, 26, 30, 36, 40])]
    for ax, (ttl, (fa, fb), (pa, pb), stops, ticks) in zip(axes, panels):
        ax.set_xlim(fa, fb)
        ax.set_ylim(-40, 2)
        ax.axvspan(pa, pb, ymin=0, ymax=(-3 + 40) / 42, color=shade, lw=0)
        ax.hlines(-3, pa, pb, colors="k", linestyles="--", lw=1.4)
        ax.text((pa + pb) / 2, -9, "通帶\n≥ −3 dB", ha="center", va="center", color=ink, fontsize=14)
        for sa, sb in stops:
            ax.axvspan(sa, sb, ymin=(-15 + 40) / 42, ymax=1, color=shade, lw=0)
            ax.hlines(-15, sa, sb, colors="k", linestyles="--", lw=1.4)
        ax.text((fa + fb) / 2, -33, "兩側阻帶 ≤ −15 dB", ha="center", va="center", color=ink, fontsize=14,
                bbox=dict(boxstyle="square,pad=0.3", fc="white", ec="none", alpha=0.9))
        ax.set_xticks(ticks)
        ax.set_yticks([0, -10, -20, -30])
        ax.set_xlabel("頻率 (GHz)")
        ax.set_title(ttl, color=ink, fontsize=16)
        ax.grid(color="#E6EBEF", lw=0.8)
        ax.set_axisbelow(True)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
        ax.set_box_aspect(0.95)
    fig.tight_layout(w_pad=3)
    bio = BytesIO()
    fig.savefig(bio, format="png", facecolor="white")
    plt.close(fig)
    bio.seek(0)
    return Image.open(bio).convert("RGB")


# ================================================================ 1 封面
s = prs.slides.add_slide(BLANK)
rect(s, 0.9, 2.2, 0.09, 2.0, BLUE)
textbox(s, 1.25, 2.2, 11.3, 0.9, [("像素化天線與濾波器 AI 設計", True)], size=40)
textbox(s, 1.25, 3.1, 11.3, 0.7, [("歷來實驗總整理：四代方法與成果，2023–2026", True)], size=26, color=BLUE)
textbox(s, 1.25, 4.45, 11.3, 0.5, ["鄒穎麒　國立中正大學　2026 年 10 月"], size=18, color=MUTED)
page_number(s)
set_notes(s, [
    "範圍：陳廷茂、錢鵬予、吳維文與我們第四代的天線與濾波器工作，2023/07 至 2026/10；起點 2023/07 出自交接紀錄與週報。",
    "前三代依各自碩士論文定稿整理；只出現在當時月報的數字會註明月報。",
    "第四代數字全部是 HFSS 模擬；四代之中只有 2026/03/24 月報報告過 5 個樣品的實作量測。",
    "這份回應日月光上次會議的要求：把之前做過的實驗整理起來，特別是對稱、七巧板、饋入強制放金屬。",
])

# ================================================================ 2 四代一覽
table_slide(
    "四代一覽", "同一個問題做了四代：用 AI 生成 25×25 像素金屬圖形，以 HFSS 模擬驗證是否通過規格。",
    [["世代", "期間", "方法", "成果"],
     ["第一代\n陳廷茂", "2023/07–\n2024/08", "生成網路加代理模型的線上學習：每次迭代只跑一次 HFSS，不需事先收集資料集",
      "雙層板與四層封裝板完成 n257 天線與濾波天線共四組模擬設計"],
     ["第二代\n錢鵬予", "2024/08–\n2025/07", "BiScaleNorm 穩定二值化訓練；週期翻轉像素的二階段生成，跳出局部解",
      "三種學習率都能收斂；第二階段改善第一階段在 29 GHz 的弱點；示範雙埠濾波器與 n260"],
     ["第三代\n吳維文", "2025/07–\n2026/06", "單階段循環排程、樣本篩選、連通 loss，並以理論公式算出的貼片當起點",
      "饋入連通率由 14.01% 提高到 61.98%，響應接近規格；期間報告 5 個樣品的實作量測"],
     ["第四代\n鄒穎麒", "2026/06–\n2026/10", "批次線上學習：從前代累積的設計出發，代理模型初篩，三台 HFSS 不停機模擬；出圖加菱形橋",
      "目前的模擬設定下，單埠天線、對稱天線、寬頻濾波器通過規格；雙埠濾波器尚未通過"],
     [("現況", True), ("2026/10", False), "第四代數字都是目前的模擬設定下的 HFSS 模擬；8/31 已交付 39 個設計，尚未打樣量測", None]],
    [1.35, 1.45, 4.75, 4.62], row_h=0.86, body_size=14.5,
    notes=[
        "出處：各代期間取自交接紀錄、週報與口試日期；碩論本身沒有寫期間。",
        "第一代碩論 2024/07；第二代 2025/06 口試；第三代 2026/06/30 口試。第三代的 5 個樣品出自 2026/03/24 月報。",
        "第三代 61.98% 是三項機制合用、平衡權重下的饋入連通率，對照組 14.01%；響應 loss 0.99，碩論描述為初步符合規格。",
        "第四代：單埠天線 +0.79 dB，加 0.10 mm 菱形橋後 +0.58 dB；對稱天線 +0.60 dB；寬頻濾波器 +0.41 dB；雙埠濾波器 −2.39 dB。都是目前的模擬設定下的值，終點都在不同機器重跑結果一致。",
        "當年前三代的目標不含方向圖與可製造性，定義不同，不構成評價；各代 loss 定義不同，不跨代比較數值。",
    ])

# ================================================================ 3 四代時間軸
fig_slide(
    "四代時間軸", "四代接力，日月光合作從第一代的天線系列開始；上方只標有書面記錄的日月光節點。",
    FIG / "timeline_simple.png",
    notes=[
        "出處：圖由 script/figs/report_generations.py 產生；期間取自交接紀錄與口試日期。",
        "2024/12 是現存最早標示日月光的簡報；2025/10 日月光要求先做天線、天線至少一軸對稱；2026/08/31 送板交付 39 個設計。",
    ])

# ================================================================ 4 評分方式：最差餘裕
fig_text_slide(
    "評分方式：最差餘裕",
    "每條規格在每個評分頻點離門檻多少 dB，取最差的一點，即 worst-case margin；正值代表全部規格通過。",
    one(MEET.glob("schem_w?.png")),
    [("怎麼算", ["每個評分頻點量 S11 比 −10 dB 低多少", "同一頻點量增益比 4 dBi 高多少", "所有規格、所有頻點中取最差的一點",
               "正值代表全部通過；負值代表最差的一點還差多少 dB"]),
     ("為什麼用它", ["單位是 dB，天線與濾波器可以用同一把尺", "第四代起使用；前三代各有自己的 loss，定義不同，不互相比較"])],
    img_w=6.9,
    notes=[
        "示意圖不是實際設計，只說明算法：綠色直線是每個評分頻點離門檻的距離，橘色是最差的一點。",
        "天線的評分頻點是 26.5–29.5 GHz 每 0.5 GHz 一點，共 7 點；S11 與增益兩條規格合計 14 個值，取最小。",
        "方向圖與可製造性是另外必須通過的條件，不併進最差餘裕。",
    ])

# ================================================================ 5 天線規格
fig_slide(
    "天線規格",
    "n257 頻段 26.5–29.5 GHz 內 S11 ≤ −10 dB、正向增益 ≥ 4 dBi；28 GHz 方向圖 ±45° 內下降不超過 3 dB，是實驗室加入的條件。",
    FIG / "gen4_ase" / "spec_antenna.png", trimmed=True,
    notes=[
        "出處：S11 與正向增益兩條規格從第一代碩論起沿用。",
        "方向圖條件由實驗室在 2026/06 加入：28 GHz 在 φ＝0° 與 90° 兩切面，±45° 內不低於峰值 −3 dB。第三代碩論曾把方向圖納入 loss 列為未來工作。",
        "圖上曲線是一個示範設計，淺色區標示規格的位置。",
        "2026/08 起依實驗室決定，單埠天線的帶外抑制交給後級濾波器，天線只負責帶內、方向圖與可製造。",
    ])

# ================================================================ 6 濾波器規格
fig_text_slide(
    "濾波器規格",
    "寬頻濾波器指通帶 26–30 GHz、在 11–40 GHz 全頻段評分、要求兩側寬阻帶的濾波器；兩者都看穿透 S21 與兩端反射 S11、S22。",
    spec_filter(),
    [("雙埠濾波器，24–32 GHz", ["通帶 25.5–30.5 GHz：S21 ≥ −3 dB", "阻帶 24–25、31–32 GHz：S21 ≤ −15 dB",
                            "26.5–29.5 GHz：S11、S22 ≤ −10 dB"]),
     ("寬頻濾波器，11–40 GHz", ["通帶 26–30 GHz：S21 ≥ −3 dB", "阻帶 11–19、36–40 GHz：S21 ≤ −15 dB",
                            "26.5–29.5 GHz：S11、S22 ≤ −10 dB"])],
    img_w=7.0,
    notes=[
        "左圖只畫 S21 的規格位置，淺色是不允許的範圍；S11、S22 的條件見右欄。",
        "雙埠濾波器規格 2026/08/12 由實驗室定案，24–32 GHz 每 0.5 GHz 一點。",
        "寬頻濾波器規格 2026/10/15 由實驗室定案，兩側阻帶位置對應 Ku 頻段與 n260 頻段。阻帶另有 −20 dB 的版本，目前 −15 dB 版已有通過的設計，−20 dB 版尚未通過，兩版的取捨請日月光決定。",
    ])

# ================================================================ 7 歷來需求整理
table_slide(
    "歷來需求整理", "以下為我們整理的歷來需求，如有遺漏或理解錯誤請指正。",
    [["時間", "來源", "內容", "回應"],
     ["2025/05/01", "會議紀錄", "濾波天線從架構上改動；新增不同截面的響應", "第二代加第二個埠做雙埠濾波器，並看 3D 場型"],
     ["2025/10/02", "會議紀錄", "先做天線、暫停濾波天線；更改 loss 計算方式；天線至少一軸對稱",
      "第三代改以天線為主、loss 只罰超出規格的部分、試 8 種翻轉；第四代做出左右對稱天線"],
     ["2026/06", "實驗室", "方向圖納入評分：28 GHz 兩切面 ±45°", "第四代單埠天線的必要條件"],
     ["2026/08", "實驗室", "單埠天線帶外交給後級濾波器；雙埠濾波器規格", "第四代單埠天線與雙埠濾波器"],
     ["2026/10", "實驗室", "對稱天線；寬頻濾波器規格", "第四代對稱天線與寬頻濾波器"],
     ["上次會議", "會議紀錄", "整理歷來實驗：對稱、七巧板、饋入強制放金屬等", "本報告"]],
    [1.45, 1.4, 4.95, 4.37], row_h=0.62, body_size=14,
    notes=[
        "出處：2025/05/01 是第二代完整報告的需求頁，記錄「日月光工程師希望：濾波天線從架構上改動、新增不同截面的響應來看」；檔案封面印 2024/05/01，日期依內容推定為 2025/05/01。",
        "2025/10/02 出自第三代 10 月進度報告標題為「日月光」的頁面，備註記錄「至少會有一軸的對稱」。",
        "實驗室一欄是實驗室依計畫進度訂定或轉達的規格。",
    ])

# ================================================================ 8–9 第一代
fig_slide(
    "第一代：生成網路與代理模型的線上學習",
    "生成網路輸出 0/1 金屬圖形送 HFSS，代理模型用累積的模擬結果邊跑邊學，把梯度傳回生成網路；圖中 MPGN 是生成網路，Gradient Estimation Network 是代理模型。",
    FIG / "thesis_gen1" / "fig3-1_design_flow.png", trimmed=True,
    notes=[
        "出處：陳廷茂碩論〈基於梯度估計網路設計貼片天線金屬圖形之研究〉圖 3-1，2024/07。",
        "代理模型訓練到與 HFSS 結果的均方誤差低於 0.0001 後，只負責把梯度傳回生成網路，每個設計仍由 HFSS 模擬。",
        "每次迭代只做一次 HFSS，不必事先收集資料集；四層板一次模擬至少 10 分鐘。",
        "這個生成網路加代理模型、邊模擬邊學的骨架，第二、三代都沿用。",
    ])

fig_table_slide(
    "第一代：四組模擬設計與規格對照",
    "雙層板與四層封裝板各做 n257 天線與濾波天線；n257 兩組只差一到三個頻點，濾波兩組帶外仍有較多頻點未達。左圖為四層板濾波天線。",
    FIG / "thesis_gen1" / "tab4-5_four_layer_filter_result.png",
    [["設計", "迭代次數", "未達處，圖上讀值"],
     ["雙層板 n257", "1,000", "29.5 GHz 增益約 3.5 dBi"],
     ["四層板 n257", "1,500", "26.5 GHz S11 約 −9.3 dB；29–29.5 GHz 增益約 3.8–3.9 dBi"],
     ["四層板濾波", "1,500", "帶內增益 3 點、帶外 S11 2 點、帶外增益 3 點"],
     ["雙層板濾波", "約 1,000", "帶內增益 4 點、帶外 S11 6 點、帶外增益 2 點"],
     [("留給後代", True), "25×25 像素化設計、生成網路加代理模型的線上學習、饋入固定金屬、只罰未達標頻點的 loss", None]],
    [1.55, 1.0, 5.22], img_w=4.1, row_h=0.72, body_size=13.5, center_cols=(1,),
    notes=[
        "出處：陳廷茂碩論表 4-3～4-6；未達處是我們依碩論原圖讀值，誤差約 ±0.3 dB。",
        "左圖是四層板濾波天線：第一層 25×25、第二層 48×22 格可調，饋入點周圍 8×8 格固定放金屬；紅線上界、淺藍下界、黃底 n257 頻段。",
        "規格：n257 天線 26.5–29.5 GHz S11 < −10 dB、增益 > 4 dBi；濾波天線另在 24–32 GHz 要求帶外高反射、低增益，訓練時以上下界表示。",
        "每換一組規格要重跑一次完整迭代。全部是 HFSS 模擬，當年沒有打樣，方向圖未列入優化。",
    ])

# ================================================================ 10–11 第二代
fig_text_slide(
    "第二代：兩項方法",
    "BiScaleNorm 讓二值化訓練不再卡住；二階段生成以週期翻轉像素跳出局部解。左圖為第二階段的結果。",
    FIG / "thesis_gen2" / "fig4-8_stage2_result.png",
    [("BiScaleNorm：穩定二值化訓練", ["金屬圖形只能是 0 或 1，訓練容易停住、依賴學習率",
                                  "正負兩側分開縮放後，三種學習率都能收斂"]),
     ("二階段生成：跳出局部解", ["第一階段先收斂到一個局部解，弱點在 29 GHz", "第二階段每 400 次翻轉少量像素，沒改善就退回歷史最佳圖",
                          "29 GHz 的反射係數與增益都比第一階段改善"])],
    img_w=6.4, size=16,
    notes=[
        "出處：錢鵬予碩論〈基於二階段梯度探索之貼片天線金屬圖形生成方法〉§4.2.1、§4.2.2，圖 4-4～4-8，2025/06。",
        "BiScaleNorm：學習率 0.003 與 0.0003 時原設計 loss 停住、BiScaleNorm 穩定收斂；0.00003 時兩者都收斂，依碩論判讀 BiScaleNorm 約第 400 次、原設計約第 1,000 次。",
        "左圖：(a) 金屬圖形、(b) 第二階段 loss、(c) 反射係數、(d) 增益，黑線為規格上下界、紅線為模擬結果。變異率起始 5%、週期 400，都是經驗值。",
        "全部是 HFSS 模擬；loss 為該論文自訂，不和其他世代比較。",
    ])

table_slide(
    "第二代：留給後代的設計元件",
    "除了兩項方法，第二代也留下饋入固定金屬、對角重疊填入與雙埠濾波結構，後續世代沿用。",
    [["項目", "第二代的做法", "後續"],
     ["饋入固定金屬", "每次生成後在饋入端疊一塊固定金屬，確保能量進入圖形", "四代都保留"],
     ["對角導通", "模擬時每格畫成 0.21 mm，對角相接的像素重疊而導通", "0.21 mm 沿用為之後的模擬格寬"],
     ["孤立金屬", "刪除會改變響應，因此保留", "後代改在生成端處理"],
     ["雙埠濾波器", "微帶線從圖形兩側穿過，看 S11、S22、S21", "第三、四代濾波器的主線"],
     ["換頻段", "同一流程改到 n260，只改 HFSS 模擬設定", "示範性結果"],
     ["代理模型", "指出代理模型有梯度誤差，列為未來方向", "後代反覆碰到的課題"]],
    [2.3, 6.2, 3.67], row_h=0.62, body_size=15,
    notes=[
        "出處：錢鵬予碩論圖 4-2、附錄一、附錄二、第五章。",
        "雙埠濾波器回應日月光 2025/05/01 的「濾波天線從架構上改動」；雙埠與 n260 各為一組示範曲線。",
    ])

# ================================================================ 12–14 第三代
fig_text_slide(
    "第三代：單階段線上學習與三項機制",
    "沿用生成器加代理模型的線上學習，加入三項機制，並以理論公式算出的 28 GHz 貼片當起點。",
    FIG / "thesis_gen3" / "fig3-11_design_flow.png",
    [("三項機制", ["ACP：訓練停滯時自動重新拉高學習率，把前代的兩階段合成一段",
                "DLF：只挑表現優於累計平均的樣本給代理模型複習",
                "SC loss：訓練中要求金屬連成一塊，抑制孤島"]),
     ("起點", ["用傳輸線公式算出 28 GHz 貼片當初始圖形，不從空白開始"])],
    img_w=7.0, size=16,
    notes=[
        "出處：吳維文碩論〈自適應循環策略與圖譜連通度損失函數於貼片天線金屬圖形生成之研究〉圖 3-11，2026/06。",
        "規格：n257 26.5–29.5 GHz，S11 < −10 dB、單一方向增益 > 4 dB。",
        "每次迭代仍跑 HFSS；代理模型提供梯度，不取代模擬。",
    ])

fig_table_slide(
    "第三代：逐項加入的比較與孤島抑制",
    "逐項加入三項機制，各列條件不同、只作定性參考；響應 loss 為 0 代表帶內全部達標，加入 SC loss 後饋入連通率由 14.01% 提高到 61.98%。",
    FIG / "thesis_gen3" / "tab4-3_base_tv_sc_radiation.png",
    [["設定", "響應 loss", "饋入連通率"],
     ["固定學習率，不加機制", "3.64", "27.22%"],
     ["只加 DLF", "1.71", "16.28%"],
     ["只加 ACP", "2.07", "18.89%"],
     ["ACP 與 DLF，左圖 Base", "0.0", "14.01%"],
     [("三項合用，左圖 SC", True), ("0.99", True), ("61.98%", True)]],
    [2.95, 1.55, 1.75], img_w=5.6, row_h=0.62, body_size=15, center_cols=(1, 2),
    notes=[
        "出處：吳維文碩論表 4-4、表 4-3、圖 4-8、圖 4-9。",
        "饋入連通率：從饋入點只走上下左右能連到的金屬占全部金屬的比率，對角相接不算連通。左圖綠色是連通的金屬、灰色是孤島，下排是 28 GHz 場型。",
        "響應 loss 只算超出規格的部分，是第三代自己的定義，不跨代比較。只加 ACP 的一列訓練 500 次、只加 DLF 的一列 1,000 次。",
        "左圖三種設定都用理論貼片起點、ACP 與 DLF，差別只在孤島抑制：不加、TV loss 22.39%、SC loss 61.98%；SC 權重過大時連通率可到 98.18%，但 S11 與增益不達標。",
    ])

fig_text_slide(
    "第三代期間的 5 個實作樣品量測",
    "2026/03/24 聯合審查報告 5 個單元天線樣品的無反射室量測；量測通帶比模擬整體往高頻移。",
    FIG / "thesis_gen3" / "ase_202603final_sample1_s11.png",
    [("回波損耗 > 10 dB 的通帶，模擬 → 量測", [
        "樣品 1：26.1–30.0 → 26.9–31.6 GHz", "樣品 2：26.2–30.1 → 26.9–31.6 GHz",
        "樣品 3：26.2–29.9 → 26.7–31.4 GHz", "樣品 4：26.1–30.2 → 26.7–31.6 GHz",
        "樣品 5：26.2–30.1 → 26.7–31.5 GHz"]),
     ("觀察", ["量測下緣 26.7–26.9 GHz，未涵蓋 n257 下緣 26.5 GHz"]),
     ("留給後代", ["饋入連通率指標、連通 loss、理論貼片起點，以及四代唯一的實測經驗"])],
    img_w=6.2, size=15,
    notes=[
        "出處：2026/03/24 ASE-CCU 聯合審查報告，屬當時月報，不在碩論內。左圖是樣品 1 的 S11，紅虛線模擬、藍線量測。",
        "模擬已含 K 接頭與 ENIG 化金層；量測在無反射室進行。由報告數字推算，量測下緣高 0.5–0.8 GHz、上緣高 1.4–1.6 GHz。",
        "28 GHz 峰值增益模擬／量測：樣品 1 至 5 為 2.2／3.0、2.0／2.7、1.3／2.7、0.9／2.1、1.5／2.4 dBi。",
        "2026/05/28 月報：顯微鏡量得實作像素 0.188–0.231 mm，模擬為 0.21 mm。樣品的設計流程與過蝕刻分析本報告未納入。",
    ])

# ================================================================ 15 第四代：接手後接續的三件事
columns_slide(
    "第四代：接手後接續的三件事",
    "前代留下生成器、代理模型與 34,215 筆模擬，第四代各任務都從前代累積的設計出發；當年目標不含方向圖與可製造性，我們接手後接續這三件事。",
    [("可製造", ["懸浮金屬與 1–3 像素的小碎片", "對角相接的像素只靠 0.01 mm 重疊相連", "依一般 PCB 規則有導通疑慮，請日月光確認",
               "直接刪碎片，響應會大幅變差"]),
     ("規格擴充", ["方向圖：28 GHz 兩切面 ±45° 內不低於峰值 −3 dB", "金屬左右對稱的天線", "雙埠濾波器與 11–40 GHz 寬頻濾波器"]),
     ("效率", ["單機逐筆：生成、模擬、訓練", "每筆模擬後等代理模型訓練", "調整算法時模擬機停下來等"])],
    size=17,
    notes=[
        "前代資料：承接的歷史模擬 34,215 筆，單埠 24,189、雙埠約 1.0 萬。單埠天線與雙埠濾波器的起點直接取自前代資料；對稱天線與寬頻濾波器的起點是由前代設計演變而來的我們較早的設計。",
        "小碎片：把前代資料中 5 個達標或接近達標的設計拔掉 1–3 像素小碎片，4 個大幅變差，下降 4.7–16.9 dB。懸浮金屬：用 1 像素把懸浮金屬接回主體，6 個試驗中有 4 個的最差餘裕掉到 −5.4 至 −5.9 dB。小碎片與懸浮金屬參與共振，不能事後直接刪或接。",
        "對角接點：模擬時每格金屬畫成 0.21 mm，對角相鄰處只重疊 0.01 mm。",
        "當年的優化目標不含方向圖與可製造性，定義不同，不構成評價。",
    ])

# ================================================================ 16 第四代：方法流程
s, top = content_slide(
    "第四代：批次線上學習流程",
    "代理模型負責大量初篩，人決定搜尋方向；研究者判讀與調整時，模擬機照常運轉。",
    notes=[
        "上圖是前代流程示意：產生一個候選、HFSS 模擬、訓練代理模型依序循環，每筆模擬後要等代理模型訓練完，研究者調整算法時模擬機也要等。",
        "下圖是批次線上學習：開發機一次挑出一組候選放進佇列，三台模擬機各自不停取件；收齊後重訓代理模型再挑下一組。所有數字以 HFSS 為準。",
        "一天約 1,600 筆是換算值：每筆中位約 160 秒，三台各約 540 筆；實測最多的一天是 1,554 筆。到 10 月自行累積的 HFSS 結果超過 7.6 萬筆，含重複模擬。",
        "每一輪開始前先寫下判斷成功與失敗的標準。示意圖不按實際時間比例。",
    ])
lg_y = top + 0.05
lx = 6.25
for col, name in ((TEAL, "挑選或產生候選"), (NAVY, "HFSS 模擬"), (ORANGE, "訓練代理模型"), (GRAY, "人工調整")):
    rect(s, lx, lg_y + 0.09, 0.26, 0.2, col)
    textbox(s, lx + 0.3, lg_y, 1.75, 0.38, [name], size=13, color=INK)
    lx += 0.3 + _text_width(name, 13) + 0.28
y1 = top + 0.5
textbox(s, L - 0.04, y1, 11, 0.4, [("前代：線上學習，1 台模擬機", True, BLUE)], size=20)
x0, lab_w = L + 1.45, 1.4
textbox(s, L, y1 + 0.55, lab_w, 0.5, ["模擬機"], size=15, anchor=MSO_ANCHOR.MIDDLE)
bx = x0
seq = [(TEAL, 0.18), (NAVY, 0.5), (ORANGE, 0.34)]
for k in range(4):
    for col, w in seq:
        rect(s, bx, y1 + 0.55, w, 0.5, col)
        bx += w
gw = (R - x0) - 8 * 1.02
rect(s, bx, y1 + 0.55, gw, 0.5, GRAY)
textbox(s, bx, y1 + 0.55, gw, 0.5, ["人工調整算法，模擬機停下等待"], size=12.5, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
bx += gw
while bx + 1.02 <= R + 0.001:
    for col, w in seq:
        rect(s, bx, y1 + 0.55, w, 0.5, col)
        bx += w
y2 = y1 + 1.4
textbox(s, L - 0.04, y2, 11.5, 0.4, [("第四代：批次線上學習，3 台模擬機幾乎不空轉，一天約 1,600 筆", True, BLUE)], size=20)
rows_y = [y2 + 0.55 + k * 0.6 for k in range(4)]
for name, yy in zip(("開發機", "HFSS 1", "HFSS 2", "HFSS 3"), rows_y):
    textbox(s, L, yy, lab_w, 0.44, [name], size=15, anchor=MSO_ANCHOR.MIDDLE)
span = R - x0
dev = [(0.0, 0.62, TEAL, "挑選"), (0.95, 1.2, ORANGE, "重訓代理模型"), (2.45, 2.0, GRAY, "人工判讀與調整"),
       (4.75, 0.62, TEAL, "挑選"), (5.7, 1.2, ORANGE, "重訓代理模型"), (7.2, 2.0, GRAY, "人工判讀與調整"),
       (9.5, 0.62, TEAL, "挑選")]
for off, w, col, txt in dev:
    w = min(w, span - off)
    rect(s, x0 + off, rows_y[0], w, 0.44, col)
    textbox(s, x0 + off, rows_y[0], w, 0.44, [txt], size=12.5, color=INK if col == GRAY else WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for k, start in enumerate((0.15, 0.32, 0.49)):
    bx = x0 + start
    while bx < R - 0.05:
        w = min(0.56, R - bx)
        rect(s, bx, rows_y[k + 1], w, 0.44, NAVY)
        bx += 0.62
ya = rows_y[-1] + 0.62
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(ya), Inches(R), Inches(ya))
ln.line.color.rgb = MUTED
ln.line.width = Pt(1.25)
lnel = ln.line._get_or_add_ln()
lnel.append(lnel.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
textbox(s, R - 1.0, ya + 0.02, 1.0, 0.35, ["時間"], size=13, color=MUTED, align=PP_ALIGN.RIGHT)

# ================================================================ 17–20 第四代：四任務爬升圖
CLIMB = FIG / "gen4_ase"
fig_slide(
    "單埠天線：最差餘裕的爬升",
    "從前代累積的設計出發：27,433 筆模擬、29 天，從 −2.68 到 +0.79 dB；加上菱形橋的可製造版本為 +0.58 dB。",
    CLIMB / "climb_single.png",
    notes=[
        "縱軸是最差餘裕，只算 26.5–29.5 GHz 的 S11 與增益；藍點是方向圖與可製造條件都通過的設計，灰點是沒通過的，累計最佳線只計藍點；低於 −4 dB 的點畫在底線上。",
        "起點 −2.68 dB：前代資料中的一個設計拔掉 1–3 像素小碎片後的版本，7/03 第一輪模擬中的最佳值。",
        "第 346 筆首次通過，+0.15 dB：部分左右對稱化讓碎片消失，再從好設計局部修改。中段 +0.49 到 +0.56 dB：大量產生好設計的變體，代理模型初篩後模擬。終點 +0.79 dB，7/31：移植另一款設計的局部結構，再局部修改；方向圖餘裕 +0.20 dB，在不同機器重跑三次結果一致。",
        "可製造版本：同一設計加 0.10 mm 菱形橋為 +0.58 dB，在不同機器重跑一致。以更嚴的求解器網格設定模擬時，這個設計是 −0.25 dB，數字不能當成量測預測。",
        "筆數只算不重複圖形，含模擬機空閒時自動產生的隨機設計，主動挑選的是 16,680 筆；期間 7/03–7/31。全部是 HFSS 模擬。",
    ])

fig_slide(
    "對稱天線：最差餘裕的爬升",
    "從既有天線設計出發，金屬精確左右鏡射：5,144 筆模擬、9 天，從 −3.53 到 +0.60 dB；方向圖未列入評分。",
    CLIMB / "climb_sym.png",
    notes=[
        "回應日月光 2025/10「天線至少一軸對稱」。生成時金屬精確左右鏡射，可自由變動的像素由 624 減為 324，對角接點全部加 0.10 mm 菱形橋。",
        "起點 −3.53 dB：以我們的單埠好設計左右鏡射當種子，第一輪 16 筆中的最佳；第二輪由代理模型挑選後升到 −0.32 dB。",
        "約第 2,240 筆首次通過，約 +0.04 dB，屬邊緣通過，前 5,000 筆只有這一筆。終點 +0.60 dB，10/15：從好設計一次翻一組左右對稱的像素；在不同機器重跑結果一致，S11 餘裕 +0.70、增益餘裕 +0.60 dB。",
        "評分只看 S11 與正向增益，方向圖未列入。10/12–10/14 沒有模擬對稱天線。10/15 之後另有只跑過一次、尚未重跑確認的較高值，不列入。全部是 HFSS 模擬。",
    ])

fig_slide(
    "雙埠濾波器：最差餘裕的爬升",
    "從前代累積的設計出發：15,625 筆模擬、12 天，從 −5.90 到 −2.39 dB，未通過規格；卡在通帶邊緣與阻帶起點，目前暫停。8/31 交付的 14 個濾波器也都未通過。",
    CLIMB / "climb_dual.png",
    notes=[
        "規格：24–32 GHz 每 0.5 GHz 一點；S11、S22 ≤ −10 dB，S21 通帶 25.5–30.5 GHz ≥ −3 dB，阻帶 24–25 與 31–32 GHz ≤ −15 dB；可製造幾何、對角接點加 0.075 mm 菱形橋。低於 −9 dB 的點畫在底線上。",
        "起點 −5.90 dB，8/13：前代雙埠設計原樣換成可製造幾何並加橋，第一輪 90 筆中的最佳。",
        "從好設計局部修改：同一天每步翻 1–2 個像素，到 −4.26、−3.82 dB。代理模型挑選、模擬機不停機補量：8/16 到 −2.88 dB，8/18 一天內刷新四次。自動化流程：腳本自動挑選與重跑驗證，最後兩次刷新 −2.43、−2.39 dB 由它產生；終點 8/24，在不同機器重跑一致。",
        "8/13 以前的舊模擬幾何把零厚度的對角點接觸當成導通，那段的數字不放在這張圖。0.075 mm 橋低於一般 PCB 最小線寬 0.1 mm，請日月光確認。",
        "筆數含空閒時的隨機設計，主動挑選 8,450 筆。全部是 HFSS 模擬。",
    ])

fig_slide(
    "寬頻濾波器：最差餘裕的爬升",
    "從既有濾波器設計出發，阻帶 −15 dB 版：2,252 筆模擬、9 天，從 −10.99 到 +0.41 dB。",
    CLIMB / "climb_wideband.png",
    notes=[
        "規格：11–40 GHz；S21 通帶 26–30 GHz ≥ −3 dB、阻帶 11–19 與 36–40 GHz ≤ −15 dB；S11、S22 在 26.5–29.5 GHz ≤ −10 dB；對角接點加 0.075 mm 菱形橋。",
        "前段以阻帶 −20 dB 版為挑選目標，圖上全部統一用 −15 dB 版計分。−20 dB 版目前最佳約 −3.05 dB，只跑過一次、尚未重跑確認，卡在 11–19 GHz 低阻帶。",
        "起點 −10.99 dB：前幾版規格下我們自己的好設計改用 11–40 GHz 重新模擬，第一輪 24 筆中的最佳；這些設計在 13.5–14 GHz 都有寄生通帶。",
        "兩埠金屬斷開、大步變異：一次改 30 格，把寄生峰壓下去，到 −5.75 dB。代理模型改學「改哪幾格、曲線怎麼變」，用來排序候選，到 −1.76 dB。好設計之間互相移植區塊，第 1,371 筆首次通過，約 +0.01 dB。通帶兩側成對修改，同時抬起 26 與 30 GHz 兩緣，到 +0.36 dB；再移植到其他好設計，終點 +0.41 dB，10/23，在不同機器重跑曲線完全一致。",
        "求解器設定與天線同等級，沒有做網格收斂檢查。全部是 HFSS 模擬。",
    ])

# ================================================================ 21 量化總表
table_slide(
    "第四代：四個任務的量化總表",
    "所有數字為目前的模擬設定下的 HFSS 模擬，模擬網格尚未做收斂驗證；寬頻濾波器以阻帶 −15 dB 版計分。",
    [["任務", "起點最差餘裕", "終點最差餘裕", "模擬筆數", "天數", "狀態"],
     ["單埠天線", "−2.68 dB", "+0.79 dB\n加橋 +0.58 dB", "27,433", "29 天", "通過\n目前的模擬設定"],
     ["對稱天線", "−3.53 dB", "+0.60 dB", "5,144", "9 天", "通過\n目前的模擬設定"],
     ["雙埠濾波器", "−5.90 dB", "−2.39 dB", "15,625", "12 天", "未通過"],
     ["寬頻濾波器", "−10.99 dB", "+0.41 dB", "2,252", "9 天", "通過\n目前的模擬設定"]],
    [2.45, 2.0, 2.3, 1.85, 1.4, 2.17], row_h=0.92, header_size=16, body_size=18,
    center_cols=(1, 2, 3, 4, 5), header_h=0.55,
    notes=[
        "口徑：最差餘裕是每條規格在每個評分頻點離門檻多少 dB、取最差一點，正值代表全部通過；模擬筆數只算不重複的圖形；起點是第一輪模擬中的最佳值；天數是第一筆到終點那一筆的日曆天數，含頭尾。",
        "期間：單埠天線 7/03–7/31；對稱天線 10/07–10/15；雙埠濾波器 8/13–8/24；寬頻濾波器 10/15–10/23。四個終點都在不同機器重跑結果一致。",
        "起點來源：單埠天線與雙埠濾波器取自前代資料；對稱天線與寬頻濾波器是我們較早的設計。模擬筆數含模擬機空閒時自動產生的隨機設計，主動挑選的是單埠 16,680、雙埠 8,450。",
        "條件：單埠天線終點無菱形橋，加 0.10 mm 橋後 +0.58 dB；對稱天線含 0.10 mm 橋；兩個濾波器含 0.075 mm 橋。以更嚴的求解器網格設定模擬時，單埠天線 +0.79 的設計變成 −0.25 dB，抽測的 3 個通過設計都轉為不通過，送板前需要複驗。",
        "寬頻濾波器阻帶 −20 dB 版尚未通過。",
    ])

# ================================================================ 22–24 最佳設計
fig_slide(
    "最佳設計：單埠天線",
    "最差餘裕 +0.79 dB、方向圖餘裕 +0.20 dB；加上 0.10 mm 菱形橋的可製造版本為 +0.58 dB。HFSS 模擬，未實測。",
    CLIMB / "best_single.png",
    notes=[
        "由左至右：金屬圖形，饋點在下緣中央；S11 與正向增益，虛線是門檻、淺紅是評分頻段；28 GHz 兩切面方向圖。",
        "2026/07/31 的設計，在不同機器重跑三次結果一致；圖為無橋幾何。以更嚴的求解器網格設定模擬時為 −0.25 dB。",
        "8/31 交付的天線是同款式中加橋後的代表設計，不是這個設計本身。",
    ])

fig_slide(
    "最佳設計：對稱天線",
    "金屬精確左右鏡射，對角接點加 0.10 mm 菱形橋，最差餘裕 +0.60 dB；方向圖未列入評分，右圖為 28 GHz 兩切面。HFSS 模擬，未實測。",
    CLIMB / "best_sym.png",
    notes=[
        "由左至右：金屬圖形，饋點在下緣中央、左右鏡射；S11 與正向增益，S11 餘裕 +0.70、增益餘裕 +0.60 dB；28 GHz 兩切面方向圖。",
        "2026/10/15 的設計，在兩台不同機器重跑，響應完全一致；這只代表數值可重現，不代表與實測相符。",
        "這個設計不在 8/31 交付內；是否加入下次送板，請日月光決定。",
    ])

fig_slide(
    "最佳設計：寬頻濾波器",
    "阻帶 −15 dB 版通過，最差餘裕 +0.41 dB，最差的一項是通帶；阻帶 −20 dB 版尚未通過。兩埠金屬不相連，HFSS 模擬，未實測。",
    CLIMB / "best_wideband.png",
    notes=[
        "由左至右：金屬圖形，上下緣中央是兩埠的接點；S21、S11、S22，虛線是阻帶 −15 dB 版的門檻、淺紅是評分頻段。",
        "各項餘裕：通帶 +0.41、低阻帶 +1.04、高阻帶 +2.95、S11 +1.24、S22 +4.71 dB；通帶 26–30 GHz 內 S21 最差一點約 −2.6 dB。",
        "同一設計以阻帶 −20 dB 版重新評分是 −3.96 dB，卡在 11–19 GHz 低阻帶。2026/10/23 的設計，在兩台不同機器重跑，曲線差為 0；另有次選設計 +0.40 dB，也在不同機器重跑一致。",
        "兩埠金屬斷開的理由：11 GHz 時兩埠金屬相連的設計 S21 約 −3 dB，幾乎直通；斷開的約 −18 dB。對角接點加 0.075 mm 菱形橋，低於一般 PCB 最小線寬 0.1 mm，請日月光確認。",
    ])

# ================================================================ 25 可製造：孤島與對角導通
fig_table_slide(
    "可製造：孤島與對角導通",
    "四代一致發現直接刪孤島會傷響應；第四代出圖時在對角接點加菱形橋，天線 0.10 mm、濾波器 0.075 mm。",
    MEET / "schem_diamond.png",
    [["世代", "做法與結果"],
     ["第一代", "口試委員指出對角相接的像素實際製作可能斷；提出四邊各擴 0.05 mm 的構想，未實作"],
     ["第二代", "模擬時每格畫成 0.21 mm，對角相接處重疊導通；刪除孤立金屬會改變響應，因此保留"],
     ["第三代", "訓練中加連通 loss，饋入連通率提高到 61.98%"],
     ["第四代", "不刪金屬，出圖加菱形橋：3,319 個天線設計加橋後中位變化 −0.07 dB，直接切斷中位下降 4.84 dB"],
     [("請日月光確認", True), "低於一般 PCB 0.1 mm 的細橋製程規則；8/31 交付中有 18 個設計的橋寬低於 0.1 mm"]],
    [1.55, 5.3], img_w=5.0, row_h=0.74, body_size=14,
    notes=[
        "左圖照實際模擬幾何繪製：原始幾何的對角接點只靠 0.01 mm 重疊相連；直接切斷是在角點挖掉一顆 0.10 mm 的 45° 菱形；加橋是在角點補一顆同尺寸的菱形金屬。示意圖不代表特定設計。",
        "2026/08/05–08/10 對 3,319 個天線設計各模擬加橋與切斷兩種幾何，每個設計各模擬一次；加橋後 86.5% 的設計變化在 0.3 dB 以內，對角接點 17 座以上時這個比例降到約四成。單埠天線 +0.79 的設計加橋後為 +0.58 dB。",
        "濾波器的舊模擬幾何把零厚度的對角點接觸當成導通，8/13 起改用與天線相同的可製造幾何；同一個濾波器設計切掉橋、0.01 mm 微橋、0.075 mm 橋分別是 −31.03、−6.86、−3.82 dB。",
        "第一代出自口委回覆文件；第二代出自碩論附錄二與口委答覆；第三代出自碩論，依序試過的 TV、IS、GC loss 出自 2025/12–2026/02 月報，IS 即碩論中的 KNS。",
    ])

# ================================================================ 26–28 日月光指定主題
s = table_slide(
    "日月光指定主題：對稱",
    "完全對稱要在生成時內建；事後完全對稱化會傷響應，局部對稱化可修掉碎片。饋入在下緣中央，所以天線只能左右對稱。",
    [["世代", "做了什麼", "結果"],
     ["第一代", "未做", "—"],
     ["第二代", "2024/09 試過對稱生成，想改善場型凹陷", "無留存紀錄、沒有延續；碩論未收錄"],
     ["第三代", "鏡射擴增資料、對稱式七巧板；2025/10 起做 8 種翻轉，先由代理模型預測再挑 1–2 個送模擬",
      "停在代理模型的預測，沒有 HFSS 數值結論；碩論未收錄"],
     ["第四代，7 月", "前代散亂設計做局部左右對稱", "碎片消失，最差餘裕由 −1.42 升到 −0.29 dB，之後首次通過規格"],
     ["第四代，8 月", "13 個好設計事後完全對稱化", "波束都置中，但最差餘裕中位下降 17.16 dB"],
     ["第四代，10 月", "生成時金屬精確左右鏡射，再左右成對局部修改", ("+0.60 dB，通過規格；方向圖未列入評分", True)]],
    [2.0, 5.3, 4.87], row_h=0.64, body_size=14.5,
    notes=[
        "出處：第二代出自研究彙整簡報；第三代出自 2025/07–2025/12 月報，圖上數字是代理模型預測值，不是 HFSS 結果；第四代出自研究日誌。",
        "第四代 7 月：−1.42 是含碎片的原件，局部左右對稱化後的無碎片版本為 −0.29 dB，在不同機器重跑一致；再從好設計局部修改，單埠天線首次通過規格。",
        "第四代 8 月：13 個已通過規格的設計做三種鏡射，共 34 筆，每個設計模擬一次。濾波器的對照顯示生成時就對稱與不對稱的品質相當，因此濾波器把生成時對稱設為預設。",
        "第四代 10 月：先普查已模擬過的 39,111 張設計，有 1,030 張完全左右對稱；生成時就鏡射的 5,000 筆中只有一筆邊緣通過，改成左右成對局部修改後到 +0.60 dB，含 0.10 mm 菱形橋。全部是 HFSS 模擬。",
    ])

table_slide(
    "日月光指定主題：七巧板",
    "七巧板指從固定形狀庫挑小形狀拼滿 25×25 圖形，原意是保證導通與可製造；只有第三代正式做過，沒有形成數值結論。",
    [["世代", "做了什麼", "結果"],
     ["第一代", "未做", "—"],
     ["第二代", "未做；相近的嘗試是每格從幾種小圖案中選一種，目的是減少像素不相連", "接上代理模型後訓練不起來，沒有延續；碩論未收錄"],
     ["第三代", "2025/07 起陸續擴充，形狀庫由 8 個粗形狀擴到 35 種，含 T、L、U 型各種旋轉；比較兩種可微的挑選方式",
      "只有曲線與圖形，沒有數值結論；2026 年起月報不再出現，碩論未收錄"],
     ["第四代", "未做；導通與可製造改用出圖加菱形橋處理", "若日月光希望把設計規則內建進生成，形狀庫值得評估"]],
    [1.7, 6.0, 4.47], row_h=0.85, body_size=14.5,
    notes=[
        "出處：第三代 2025/09–2025/12 月報；月報記錄過程與圖形，沒有寫數值結論與停止的理由。第二代出自研究彙整簡報。",
        "第三代各版：2025/07 用 8 個 5×5 粗形狀；2025/10 加入細線、十字等，14、16 個形狀；2025/11 擴成 35 種形狀庫，並比較 Gumbel-Softmax 與 Gumbel-Sinkhorn 兩種可微的挑選方式。",
        "第四代沒有做七巧板；相近的做法是在好設計之間移植整塊金屬區塊，但不是從固定形狀庫拼滿。",
    ])

table_slide(
    "日月光指定主題：饋入強制放金屬",
    "四代都固定饋入金屬，這是確保饋入接觸的前提；效果只有第一代的一次對照，沒有做過系統比較。寬頻濾波器另發現兩埠之間的金屬要斷開。",
    [["世代", "做了什麼", "結果"],
     ["第一代", "碩論實驗一：饋入點周圍 8×8 格強制放金屬", "固定後 loss 較快、較穩地下降，之後都沿用"],
     ["第二代", "每次生成後在饋入端疊一塊固定金屬；雙埠改成上下各一塊", "碩論正式寫入，一直沿用"],
     ["第三代", "比較有無固定；2026/05 固定區由 5×5 擴為 9×9",
      "一度觀察到不固定較好，但當時只在變異後才套上；9×9 時饋入連通率約 52%"],
     ["第四代", "單埠：饋入點永遠是金屬，必要時在下緣中央補 3×3 格金屬；雙埠：上下緣中央各固定 5×5 格金屬當接點",
      "全程固定；寬頻濾波器兩埠金屬斷開時 11 GHz 穿透約 −18 dB，相連時約 −3 dB"]],
    [1.6, 5.9, 4.67], row_h=0.88, body_size=14,
    notes=[
        "出處：第一代為陳廷茂碩論實驗一，兩組各跑一次，論文沒有給數字；第二代為錢鵬予碩論圖 4-2；第三代為 2025/09 與 2026/05/28 月報，9×9 時饋入連通率 52.42%；第四代出自研究日誌。",
        "寬頻濾波器的發現：11–19 GHz 低阻帶要壓到 −15 dB 以下，實際上需要兩埠金屬斷開；這是固定饋入金屬之外的新條件。全部是 HFSS 模擬。",
    ])

# ================================================================ 29 現況與下一步
table_slide(
    "現況與下一步",
    "截至 2026/10/23，所有第四代數字都是 HFSS 模擬；右兩欄是可討論的選項與請日月光決定的事，尚未決定。",
    [["項目", "現況", "可討論的選項", "請日月光決定"],
     ["單埠天線", "通過，+0.79 dB；加 0.10 mm 橋 +0.58 dB", "打樣量測", "—"],
     ["對稱天線", "通過，+0.60 dB，含 0.10 mm 橋", "打樣量測", "是否加入下次送板"],
     ["雙埠濾波器", "未通過，−2.39 dB，目前暫停", "依需求再啟動", "—"],
     ["寬頻濾波器", "阻帶 −15 dB 版通過，+0.41 dB；−20 dB 版未通過，最佳約 −3.05 dB", "送板打樣，或繼續推進低阻帶",
      "阻帶要 −15 還是 −20 dB"],
     ["8/31 交付", "39 個設計：天線 25、濾波器 14，尚未實測；濾波器全部未通過規格", "打樣量測", "細橋製程規則"],
     ["模擬限制", "較嚴的模擬設定下，抽測 3 個通過的設計全部轉為不通過，單埠 +0.79 dB 變成 −0.25 dB；送板前需要複驗",
      "送板前以較嚴設定複驗", "—"],
     ["實測偏移", "第三代實測通帶下緣比模擬高約 0.5–0.8 GHz；第四代設計尚未納入此偏移", "量測後回饋設計", "—"]],
    [1.75, 4.55, 2.95, 2.92], row_h=0.6, body_size=13.5,
    notes=[
        "模擬限制：所有第四代數字使用同一套 HFSS 求解器設定，尚未做網格收斂。以更嚴的網格設定模擬時共振頻率會移動，單埠天線 +0.79 的設計變成 −0.25 dB，抽測的 3 個通過設計都轉為不通過。",
        "實測偏移：第三代 5 個樣品的量測下緣比模擬高 0.5–0.8 GHz、上緣高 1.4–1.6 GHz，出自 2026/03/24 月報。第四代的設計還沒有實測，模擬與量測會差多少目前沒有數據。",
        "8/31 交付：天線 25 個，含 7 個款式的代表 21 個與前代資料中的頂尖設計 4 個；濾波器 14 個主要卡在 S21 通帶邊緣與阻帶起點。39 個設計全部在不同機器重跑，37 個完全一致、2 個差 0.08 dB 以內；18 個設計的橋寬低於一般 PCB 最小線寬 0.1 mm。",
        "寬頻濾波器 −3.05 dB 只跑過一次、尚未重跑確認；若維持 −20 dB，需要新的兩埠斷開款式當起點。對稱天線 10/15 之後另有只跑過一次的較高值，不列入。",
        "對稱天線與寬頻濾波器都是 8/31 之後的結果，不在已交付的 39 個設計內。以上都是選項，供討論，不代表已決定。",
    ])

# ================================================================ 30 附錄：名詞白話表
s, top = content_slide("附錄：名詞白話表", "第一次讀到不懂的詞先查這裡。", notes=[
    "名詞以本報告的用法為準；各代方法名稱沿用碩論原名。"])
TERMS_L = [
    ("像素化金屬圖形", "25×25 格、格距 0.2 mm，逐格填金屬或留空"),
    ("0.21 mm 重疊填入", "模擬時每格畫成 0.21 mm，對角處重疊 0.01 mm"),
    ("HFSS", "電磁全波模擬軟體；四代都用它驗證設計"),
    ("目前的模擬設定", "第四代共用的求解器設定；更嚴時數字會移動"),
    ("最差餘裕", "各規格在各頻點離門檻多少 dB，取最差一點"),
    ("方向圖", "各角度的輻射強度；本計畫看 28 GHz 兩切面"),
    ("代理模型", "用神經網路近似 HFSS，用來傳梯度或初篩"),
    ("線上學習", "生成、模擬、訓練交替進行，不需事先收資料"),
    ("批次線上學習", "一組候選送三台 HFSS，收齊後重訓再挑下一組"),
    ("局部修改", "從好設計出發，每次只改少數像素"),
    ("區塊移植", "把一個好設計的一塊金屬區域換到另一個設計"),
]
TERMS_R = [
    ("饋入固定金屬", "饋入點附近強制放金屬，確保能量進入圖形"),
    ("埠墊", "雙埠濾波器上下緣中央固定的 5×5 格金屬"),
    ("孤島、懸浮金屬", "沒有連到饋入點的金屬"),
    ("饋入連通率", "與饋入點連成一塊的金屬占全部金屬的比率"),
    ("菱形橋", "對角接點補一顆 45° 菱形金屬，方便製作"),
    ("上下界 loss", "每個頻點只罰超出上界或低於下界的部分"),
    ("BiScaleNorm", "第二代：正負兩側分開縮放，緩解梯度消失"),
    ("二值化溫度", "控制輸出接近 0 或 1 的程度，高時便於探索"),
    ("ACP", "第三代：停滯時自動拉高學習率與二值化溫度"),
    ("DLF", "第三代：只用優於累計平均的樣本訓練代理模型"),
    ("SC loss", "第三代的連通 loss：訓練中要求金屬連成一塊"),
    ("TV、IS、GC loss", "第三代試過的其他孤島抑制 loss"),
]
tw = (CW - 0.35) / 2
for k, terms in enumerate((TERMS_L, TERMS_R)):
    table(s, L + k * (tw + 0.35), top + 0.05, [["名詞", "白話"]] + [[a, b] for a, b in terms], [1.75, tw - 1.75],
          header_size=13.5, body_size=12.5, row_h=0.395, header_h=0.38)

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
