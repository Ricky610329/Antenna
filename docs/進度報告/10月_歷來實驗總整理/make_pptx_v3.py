# -*- coding: utf-8 -*-
"""給日月光的「歷來實驗總整理」第三版投影片產生器，41 頁；第二版 make_pptx_v2.py 保留不動。

和第二版的差別：主軸改成每一代的嘗試歷程，各代依序是方法骨架、嘗試歷程、成果與留下的東西；
刪掉需求整理頁與三頁指定主題，對稱、七巧板、饋入固定金屬、孤島與導通分散寫進各代段落。
前三代的嘗試歷程頁取自 notes/genN_嘗試歷程.md，圖用 figs/attempts_genN/ 與 figs/thesis_genN/ 的學長原圖；
第四代加嘗試歷程三頁、目前架構圖一頁、新版代理模型兩頁，素材在 notes/gen4_嘗試歷程.md 與
notes/gen4_架構與代理模型.md；方向圖規格一律寫成「不低於 0° 方向增益 −3 dB」。

版式沿用第二版元件：16:9、白底、30pt 粗體標題、標題下灰細線加短深藍線、標題下一句描述、
小標 20pt 粗體深藍、表格淺藍表頭；不放底部結語帶、不放圖下小字，條件與出處寫進備忘稿。

跑法：在 repo 根執行 python docs/進度報告/10月_歷來實驗總整理/make_pptx_v3.py
→ 同夾 歷來實驗總整理_v3.pptx；本夾在 .gitignore，不進 git。
爬升圖與最佳設計卡由 script/figs/report_ase_climb.py 產生，架構與代理模型圖由 script/figs/report_ase_arch.py
產生，時間軸圖由 script/figs/report_generations.py 產生。
"""
import math
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
FIG = HERE / "figs"
MEET = REPO / "docs" / "進度報告" / "組會_2026-10" / "figs"
OUT = HERE / "歷來實驗總整理_v3.pptx"

FOOTER_BAND = False     # 底部結語帶；一律不放
SOURCE_LINE = False     # 頁底出處小字；一律不放，出處寫進備忘稿

FONT = "Microsoft JhengHei"
INK = RGBColor(0x15, 0x2C, 0x3D)
BLUE = RGBColor(0x24, 0x5B, 0x85)
MUTED = RGBColor(0x52, 0x6A, 0x7D)
RULE = RGBColor(0xCC, 0xD8, 0xE2)
HEAD = RGBColor(0xEE, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x3A, 0x8A, 0x7C)
NAVY = RGBColor(0x24, 0x4C, 0x6A)
ORANGE = RGBColor(0xC0, 0x62, 0x2F)
GRAY = RGBColor(0xCB, 0xD3, 0xDA)
W, H = 13.333, 7.5
L, R = 0.583, 12.75
CW = R - L
BODY_BOTTOM = 6.88

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def repo(rel):
    return REPO / rel


def one(paths):
    """glob 結果必須恰好一個檔。檔名含研究紀錄代號的圖用 glob 取，原始碼裡不出現代號。"""
    paths = sorted(paths)
    assert len(paths) == 1, paths
    return paths[0]


# ---------------------------------------------------------------- 基本元件，沿用初版 make_pptx.py
def _run(par, text, size, bold=False, color=INK):
    r = par.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    rpr = r._r.get_or_add_rPr()                      # 中文字也指定正黑體，避免 PowerPoint 換成佈景字型
    rpr.set("lang", "zh-TW")                         # 繁中斷行規則：右括號不放行首
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)
    return r


def textbox(slide, x, y, w, h, lines, size=19, color=INK, align=None, anchor=None, spacing=1.15, after=0):
    """lines: list of str 或 (str, bold) 或 (str, bold, color)。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.line_spacing = spacing
        if after:
            par.space_after = Pt(after)
        if align is not None:
            par.alignment = align
        item = line if isinstance(line, tuple) else (line, False)
        _run(par, item[0], size, item[1], item[2] if len(item) > 2 else color)
    return tb


def rich_text(slide, x, y, w, h, blocks, size=16, head_size=20, anchor=MSO_ANCHOR.MIDDLE):
    """右欄：blocks＝[(藍色小標, [條列...]), ...]；小標 20pt 粗體深藍，條列前加「・」。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = anchor
    first = True
    for head, items in blocks:
        if head:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            par.line_spacing = 1.1
            par.space_after = Pt(4)
            if par is not tf.paragraphs[0]:
                par.space_before = Pt(14)
            _run(par, head, head_size, True, BLUE)
        for it in items:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            par.line_spacing = 1.12
            par.space_after = Pt(6)
            ppr = par._p.get_or_add_pPr()                # 懸掛縮排：換行後對齊「・」後的文字
            ppr.set("marL", str(int(Pt(size))))
            ppr.set("indent", str(-int(Pt(size))))
            _run(par, "・" + it, size, False, INK)
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    ln = s.line._get_or_add_ln()                     # 不畫外框
    ln.append(ln.makeelement(qn("a:noFill"), {}))
    s.shadow.inherit = False
    return s


def _as_stream(img):
    if isinstance(img, Image.Image):
        bio = BytesIO()
        img.save(bio, "PNG")
        bio.seek(0)
        return bio, img.size
    return str(img), Image.open(img).size


def trim(img, thresh=242, pad=10):
    """裁掉圖四周的白邊，讓圖在頁面上放得更大；只裁邊，不動內容。"""
    im = (img if isinstance(img, Image.Image) else Image.open(img)).convert("RGB")
    gray = im.convert("L").point(lambda v: 255 if v < thresh else 0)
    box = gray.getbbox()
    if not box:
        return im
    x0, y0, x1, y1 = box
    return im.crop((max(0, x0 - pad), max(0, y0 - pad), min(im.width, x1 + pad), min(im.height, y1 + pad)))


def picture_fit(slide, img, x, y, w, h, valign="center", trimmed=False):
    """把圖放進 (x,y,w,h) 框內、保持比例、水平置中；img 可以是路徑或 PIL Image。"""
    if trimmed:
        img = trim(img)
    src, (iw, ih) = _as_stream(img)
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    top = y if valign == "top" else y + (h - ph) / 2
    return slide.shapes.add_picture(src, Inches(x + (w - pw) / 2), Inches(top), Inches(pw), Inches(ph))


def page_number(slide):
    n = len(prs.slides)
    textbox(slide, 12.15, 7.06, 0.8, 0.32, [str(n)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def set_notes(slide, notes):
    if not notes:
        return
    text = "\n".join(notes) if isinstance(notes, (list, tuple)) else notes
    slide.notes_slide.notes_text_frame.text = text


def _text_width(s, pt):
    """粗估字串寬度，單位吋：全形字 1 em、半形約 0.55 em。"""
    return sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in s) * pt / 72


def content_slide(title, desc=None, notes=None):
    """回傳 (slide, 內容區頂端 y)。描述超過一行時內容區自動下移。"""
    s = prs.slides.add_slide(BLANK)
    textbox(s, L - 0.04, 0.30, CW, 0.70, [(title, True)], size=30, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, L, 1.07, CW, 0.012, RULE)
    rect(s, L, 1.055, 1.36, 0.04, BLUE)
    top = 1.75
    if desc:
        nline = math.ceil(_text_width(desc, 17) / (CW - 0.2))
        textbox(s, L - 0.04, 1.16, CW, 0.40 * nline + 0.1, [desc], size=17, color=INK)
        top = 1.80 if nline <= 1 else (2.12 if nline == 2 else 2.45)
    page_number(s)
    set_notes(s, notes)
    return s, top


def _border(cell, color=None, width_pt=0.75):
    """只畫下框線；左右上清空。框線元素要排在填色元素之前，這是 DrawingML 規定的順序。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(qn(tag), {"w": str(int(width_pt * 12700)), "cap": "flat", "cmpd": "sng",
                                        "algn": "ctr"})
        if tag == "a:lnB" and color is not None:
            fill = ln.makeelement(qn("a:solidFill"), {})
            clr = fill.makeelement(qn("a:srgbClr"), {"val": str(color)})
            fill.append(clr)
            ln.append(fill)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.insert(i, ln)


def table(slide, x, y, rows, col_w, header_size=15, body_size=14, row_h=0.5, bold_first_col=True, header_h=None,
          center_cols=()):
    """rows[0] 是表頭；儲存格給 str 或 (str, bold)；給 None 表示與左邊儲存格合併。"""
    n_r, n_c = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * n_r))
    t = shp.table
    t.horz_banding = False
    t.first_row = True
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches((header_h or min(row_h, 0.5)) if i == 0 else row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            _border(cell, RULE if i < n_r - 1 else BLUE, 0.75 if i else 1.0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD if i == 0 else WHITE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if val is None:
                continue
            tf = cell.text_frame
            tf.word_wrap = True
            text = val[0] if isinstance(val, tuple) else val
            bold = i == 0 or (isinstance(val, tuple) and val[1]) or (bold_first_col and j == 0)
            for k, part in enumerate(text.split("\n")):
                par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                if k == 0:
                    par.text = ""
                par.line_spacing = 1.0
                if j in center_cols and not (j + 1 < n_c and row[j + 1] is None):
                    par.alignment = PP_ALIGN.CENTER
                _run(par, part, header_size if i == 0 else body_size, bold, INK)
        # 合併：連續的 None 併進左邊第一個有值的儲存格
        j = 0
        while j < n_c:
            if row[j] is not None:
                k = j
                while k + 1 < n_c and row[k + 1] is None:
                    k += 1
                if k > j:
                    t.cell(i, j).merge(t.cell(i, k))
                j = k + 1
            else:
                j += 1
    return shp


# ---------------------------------------------------------------- 頁型
def table_slide(title, desc, rows, col_w, notes=None, **kw):
    s, top = content_slide(title, desc, notes)
    table(s, L, top + 0.08, rows, col_w, **kw)
    return s


def fig_slide(title, desc, img, notes=None, pad=0.0, trimmed=False):
    s, top = content_slide(title, desc, notes)
    picture_fit(s, img, L + pad, top, CW - 2 * pad, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    return s


def fig_text_slide(title, desc, img, blocks, notes=None, img_w=6.6, size=16, trimmed=True):
    """左圖右文；blocks＝[(藍色小標, [條列...]), ...]。圖預設先裁白邊。"""
    s, top = content_slide(title, desc, notes)
    if isinstance(img, Image.Image):
        trimmed = False
    picture_fit(s, img, L, top, img_w, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    tx = L + img_w + 0.35
    rich_text(s, tx, top, R - tx, BODY_BOTTOM - top, blocks, size=size)
    return s


def columns_slide(title, desc, cols, notes=None, size=17):
    """cols: [(小標, [條列...]), ...]，兩到三欄文字。"""
    s, top = content_slide(title, desc, notes)
    n = len(cols)
    gap = 0.45
    cw = (CW - gap * (n - 1)) / n
    for k, (head, items) in enumerate(cols):
        x = L + k * (cw + gap)
        rect(s, x, top + 0.25, cw, 0.05, BLUE)
        rich_text(s, x, top + 0.4, cw, BODY_BOTTOM - top - 0.6, [(head, items)], size=size,
                  anchor=MSO_ANCHOR.TOP)
    return s


def fig_table_slide(title, desc, img, rows, col_w, notes=None, img_w=5.0, trimmed=True, **kw):
    """左圖右表：用在合併頁，圖與表各自精簡。"""
    s, top = content_slide(title, desc, notes)
    picture_fit(s, img, L, top, img_w, BODY_BOTTOM - top, valign="center", trimmed=trimmed)
    shp = table(s, L + img_w + 0.3, top + 0.1, rows, col_w, **kw)
    return s, shp


# ---------------------------------------------------------------- 濾波器規格遮罩
CJK = "C:/Windows/Fonts/msjh.ttc"


def spec_filter():
    """濾波器規格遮罩：只畫 S21 的規格位置，淺色是不允許的範圍；不放示範曲線。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    font_manager.fontManager.addfont(CJK)
    plt.rcParams.update({"font.family": [font_manager.FontProperties(fname=CJK).get_name(), "DejaVu Sans"], "font.size": 15,
                         "axes.edgecolor": "#8A9AA8", "axes.labelcolor": "#152C3D", "xtick.color": "#152C3D",
                         "ytick.color": "#152C3D"})
    shade, ink = "#F3DCD3", "#152C3D"
    fig, axes = plt.subplots(1, 2, figsize=(10.4, 5.0), dpi=220)
    panels = [("雙埠濾波器 S21 (dB)", (24, 32), (25.5, 30.5), [(24, 25), (31, 32)], [24, 26, 28, 30, 32]),
              ("寬頻濾波器 S21 (dB)", (11, 40), (26, 30), [(11, 19), (36, 40)], [11, 15, 20, 26, 30, 36, 40])]
    for ax, (ttl, (fa, fb), (pa, pb), stops, ticks) in zip(axes, panels):
        ax.set_xlim(fa, fb)
        ax.set_ylim(-40, 2)
        ax.axvspan(pa, pb, ymin=0, ymax=(-3 + 40) / 42, color=shade, lw=0)
        ax.hlines(-3, pa, pb, colors="k", linestyles="--", lw=1.4)
        ax.text((pa + pb) / 2, -9, "通帶\n≥ −3 dB", ha="center", va="center", color=ink, fontsize=14)
        for sa, sb in stops:
            ax.axvspan(sa, sb, ymin=(-15 + 40) / 42, ymax=1, color=shade, lw=0)
            ax.hlines(-15, sa, sb, colors="k", linestyles="--", lw=1.4)
        ax.text((fa + fb) / 2, -33, "兩側阻帶 ≤ −15 dB", ha="center", va="center", color=ink, fontsize=14,
                bbox=dict(boxstyle="square,pad=0.3", fc="white", ec="none", alpha=0.9))
        ax.set_xticks(ticks)
        ax.set_yticks([0, -10, -20, -30])
        ax.set_xlabel("頻率 (GHz)")
        ax.set_title(ttl, color=ink, fontsize=16)
        ax.grid(color="#E6EBEF", lw=0.8)
        ax.set_axisbelow(True)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
        ax.set_box_aspect(0.95)
    fig.tight_layout(w_pad=3)
    bio = BytesIO()
    fig.savefig(bio, format="png", facecolor="white")
    plt.close(fig)
    bio.seek(0)
    return Image.open(bio).convert("RGB")



# ---------------------------------------------------------------- 第三版新增：嘗試歷程頁型與小工具
ATT_HEAD = ["時間", "嘗試", "觀察", "後續"]


def attempts_slide(title, desc, rows, notes, col_w=(1.75, 4.25, 3.6, 2.57), body=15, row_h=0.72, head=None):
    """嘗試歷程頁：全寬四欄表，每頁最多 6 列。"""
    return table_slide(title, desc, [head or ATT_HEAD] + rows, list(col_w), notes=notes, row_h=row_h,
                       body_size=body, header_size=16)


def attempts_fig_slide(title, desc, img, rows, notes, img_w=4.7, split=(0.18, 0.36, 0.25, 0.21), body=13.5,
                       row_h=0.7, head=None, trimmed=True):
    """嘗試歷程頁：左圖右表。split＝四欄占表寬的比例。"""
    tw = CW - img_w - 0.3
    col_w = [tw * f for f in split]
    s, shp = fig_table_slide(title, desc, img, [head or ATT_HEAD] + rows, col_w, notes=notes, img_w=img_w,
                             trimmed=trimmed, row_h=row_h, body_size=body, header_size=14.5, header_h=0.45)
    return s


def pair_patterns(paths, labels, gap=90, label_px=150):
    """兩張學長原圖並排、底下加中文標籤，標籤是我們加的；原圖只裁白邊、等高縮放，不改內容。"""
    from PIL import ImageDraw, ImageFont
    ims = [trim(p) for p in paths]
    h = min(im.height for im in ims)
    ims = [im.resize((round(im.width * h / im.height), h), Image.LANCZOS) for im in ims]
    font = ImageFont.truetype(CJK, label_px)
    lab_h = int(label_px * 1.9)
    out = Image.new("RGB", (sum(im.width for im in ims) + gap * (len(ims) - 1), h + lab_h), "white")
    d = ImageDraw.Draw(out)
    x = 0
    for im, lab in zip(ims, labels):
        out.paste(im, (x, 0))
        tw = d.textlength(lab, font=font)
        d.text((x + (im.width - tw) / 2, h + label_px * 0.45), lab, font=font, fill=(0x15, 0x2C, 0x3D))
        x += im.width + gap
    return out


A1, A2, A3 = (FIG / f"attempts_gen{k}" for k in (1, 2, 3))
T1, T2, T3 = (FIG / f"thesis_gen{k}" for k in (1, 2, 3))
G4 = FIG / "gen4_ase"

# ================================================================ 開場：1 封面
s = prs.slides.add_slide(BLANK)
rect(s, 0.9, 2.2, 0.09, 2.0, BLUE)
textbox(s, 1.25, 2.2, 11.3, 0.9, [("像素化天線與濾波器 AI 設計", True)], size=40)
textbox(s, 1.25, 3.1, 11.3, 0.7, [("歷來實驗總整理：四代的嘗試歷程與成果，2023–2026", True)], size=26, color=BLUE)
textbox(s, 1.25, 4.45, 11.3, 0.5, ["鄒穎麒　國立中正大學　2026 年 10 月"], size=18, color=MUTED)
page_number(s)
set_notes(s, [
    "範圍：陳廷茂、錢鵬予、吳維文與我們第四代的天線與濾波器工作，2023/07 至 2026/10；起點 2023/07 出自交接紀錄與週報。",
    "主軸是每一代做過哪些嘗試、看到什麼、後來往哪走；對稱、七巧板、饋入強制放金屬、孤島與導通這些主題，分散寫在各代自己的段落。",
    "前三代依各自的週報、月報與碩士論文定稿整理；數字以碩論為準，只出現在當時月報的數字會註明月報。",
    "第四代數字全部是 HFSS 模擬；四代之中只有 2026/03/24 月報報告過 5 個樣品的實作量測。",
])

# ================================================================ 2 四代一覽
table_slide(
    "四代一覽", "同一個問題做了四代：用 AI 生成 25×25 像素金屬圖形，以 HFSS 模擬驗證是否通過規格。",
    [["世代", "期間", "方法", "成果"],
     ["第一代\n陳廷茂", "2023/07–\n2024/08", "生成網路加代理模型的線上學習：每次迭代只跑一次 HFSS，不需事先收集資料集",
      "雙層板與四層封裝板完成 n257 天線與濾波天線共四組模擬設計"],
     ["第二代\n錢鵬予", "2024/08–\n2025/07", "BiScaleNorm 穩定二值化訓練；週期翻轉像素的二階段生成，跳出局部解",
      "三種學習率都能收斂；第二階段改善第一階段在 29 GHz 的弱點；示範雙埠濾波器與 n260"],
     ["第三代\n吳維文", "2025/07–\n2026/06", "單階段循環排程、樣本篩選、連通 loss，並以理論公式算出的貼片當起點",
      "饋入連通率由 14.01% 提高到 61.98%，響應接近規格；期間報告 5 個樣品的實作量測"],
     ["第四代\n鄒穎麒", "2026/06–\n2026/10", "批次線上學習：從前代累積的設計出發，代理模型初篩，三台 HFSS 不停機模擬；出圖加菱形橋",
      "目前的模擬設定下，單埠天線、對稱天線、寬頻濾波器通過規格；雙埠濾波器尚未通過"],
     [("現況", True), ("2026/10", False), "第四代數字都是目前的模擬設定下的 HFSS 模擬；8/31 已交付 39 個設計，尚未打樣量測", None]],
    [1.35, 1.45, 4.75, 4.62], row_h=0.86, body_size=14.5,
    notes=[
        "出處：各代期間取自交接紀錄、週報與口試日期；碩論本身沒有寫期間。",
        "第一代碩論 2024/07；第二代 2025/06 口試；第三代 2026/06/30 口試。第三代的 5 個樣品出自 2026/03/24 月報。",
        "第三代 61.98% 是三項機制合用、平衡權重下的饋入連通率，對照組 14.01%；響應 loss 0.99，碩論描述為初步符合規格。",
        "第四代：單埠天線 +0.79 dB，加 0.10 mm 菱形橋後 +0.58 dB；對稱天線 +0.60 dB；寬頻濾波器 +0.41 dB；雙埠濾波器 −2.39 dB。都是目前的模擬設定下的值，終點都在不同機器重跑結果一致。",
        "當年前三代的目標不含方向圖與可製造性，定義不同，不構成評價；各代 loss 定義不同，不跨代比較數值。",
    ])

# ================================================================ 3 四代時間軸
fig_slide(
    "四代時間軸", "四代接力，日月光合作從第一代的天線系列開始。",
    FIG / "timeline_simple.png",
    notes=[
        "出處：圖由 script/figs/report_generations.py 產生；期間取自交接紀錄與口試日期。",
        "每一代的段落依序是方法骨架、嘗試歷程、成果與留給後代的東西。",
    ])

# ================================================================ 4 評分方式：最差餘裕
fig_text_slide(
    "評分方式：最差餘裕",
    "每條規格在每個評分頻點離門檻多少 dB，取最差的一點，即 worst-case margin；正值代表全部規格通過。",
    one(MEET.glob("schem_w?.png")),
    [("怎麼算", ["每個評分頻點量 S11 比 −10 dB 低多少", "同一頻點量增益比 4 dBi 高多少", "所有規格、所有頻點中取最差的一點",
               "正值代表全部通過；負值代表最差的一點還差多少 dB"]),
     ("為什麼用它", ["單位是 dB，天線與濾波器可以用同一把尺", "第四代起使用；前三代各有自己的 loss，定義不同，不互相比較"])],
    img_w=6.9,
    notes=[
        "示意圖不是實際設計，只說明算法：綠色直線是每個評分頻點離門檻的距離，橘色是最差的一點。",
        "天線的評分頻點是 26.5–29.5 GHz 每 0.5 GHz 一點，共 7 點；S11 與增益兩條規格合計 14 個值，取最小。",
        "方向圖與可製造性是另外必須通過的條件，不併進最差餘裕。",
    ])

# ================================================================ 5 天線規格
fig_slide(
    "天線規格",
    "n257 頻段 26.5–29.5 GHz 內 S11 ≤ −10 dB、正向增益 ≥ 4 dBi；28 GHz 方向圖在兩切面 ±45° 內不低於 0° 方向增益 −3 dB，是實驗室加入的條件。",
    G4 / "spec_antenna.png", trimmed=True,
    notes=[
        "出處：S11 與正向增益兩條規格從第一代碩論起沿用。",
        "方向圖條件由實驗室在 2026/06 加入：28 GHz 在 φ＝0° 與 90° 兩切面，±45° 內每個角度的增益都不低於 0° 方向的增益減 3 dB；基準是 0° 方向，不是整圈的最大值。第三代碩論曾把方向圖納入 loss 列為未來工作。",
        "圖上曲線是一個示範設計，淺色區標示規格的位置。",
        "2026/08 起依實驗室決定，單埠天線的帶外抑制交給後級濾波器，天線只負責帶內、方向圖與可製造。",
    ])

# ================================================================ 6 濾波器規格
fig_text_slide(
    "濾波器規格",
    "寬頻濾波器指通帶 26–30 GHz、在 11–40 GHz 全頻段評分、要求兩側寬阻帶的濾波器；兩者都看穿透 S21 與兩端反射 S11、S22。",
    spec_filter(),
    [("雙埠濾波器，24–32 GHz", ["通帶 25.5–30.5 GHz：S21 ≥ −3 dB", "阻帶 24–25、31–32 GHz：S21 ≤ −15 dB",
                            "26.5–29.5 GHz：S11、S22 ≤ −10 dB"]),
     ("寬頻濾波器，11–40 GHz", ["通帶 26–30 GHz：S21 ≥ −3 dB", "阻帶 11–19、36–40 GHz：S21 ≤ −15 dB",
                            "26.5–29.5 GHz：S11、S22 ≤ −10 dB"])],
    img_w=7.0,
    notes=[
        "左圖只畫 S21 的規格位置，淺色是不允許的範圍；S11、S22 的條件見右欄。",
        "雙埠濾波器規格 2026/08/12 由實驗室定案，24–32 GHz 每 0.5 GHz 一點。",
        "寬頻濾波器規格 2026/10/15 由實驗室定案，兩側阻帶位置對應 Ku 頻段與 n260 頻段。阻帶另有 −20 dB 的版本，目前 −15 dB 版已有通過的設計，−20 dB 版尚未通過，兩版的取捨請日月光決定。",
    ])

# ================================================================ 第一代：7 方法骨架
fig_slide(
    "第一代：生成網路與代理模型的線上學習",
    "生成網路輸出 0/1 金屬圖形送 HFSS，代理模型用累積的模擬結果邊跑邊學，把梯度傳回生成網路；圖中 MPGN 是生成網路，Gradient Estimation Network 是代理模型。",
    T1 / "fig3-1_design_flow.png", trimmed=True,
    notes=[
        "出處：陳廷茂碩論〈基於梯度估計網路設計貼片天線金屬圖形之研究〉圖 3-1，2024/07。",
        "代理模型訓練到與 HFSS 結果的均方誤差低於 0.0001 後，只負責把梯度傳回生成網路，每個設計仍由 HFSS 模擬。",
        "每次迭代只做一次 HFSS，不必事先收集資料集；四層板一次模擬至少 10 分鐘。",
        "這個生成網路加代理模型、邊模擬邊學的骨架，第二、三代都沿用；接下來三頁是第一代走到這個骨架前後的嘗試。",
    ])

# ================================================================ 8 第一代嘗試歷程：出圖方法
attempts_fig_slide(
    "第一代嘗試歷程：出圖方法",
    "HFSS 不提供梯度，第一代先後試了直接反傳、強化學習與合成梯度，2023 年 11 月定案為代理模型加資料暫存；左圖為定案後的 loss。",
    A1 / "a3_gen_with_buffer_loss_2023-11.png",
    [["2023/07–10", "全連接網路直接出圖，loss 經 HFSS 反傳", "HFSS 不提供梯度，網路無法更新", "改找繞過 HFSS 的方法"],
     ["2023/10", "強化學習：策略梯度、DDPG、Dueling DQN、PPO", "每輪需要大量模擬，獎勵起伏大", "暫停，轉向梯度法"],
     ["2023/10", "以已知可達標的 4.7 GHz 天線當驗證題", "題目確定有解，便於檢查算法", "11 月回到 28 GHz"],
     ["2023/10", "合成梯度：另訓一個網路估計梯度", "想法成形", "發展成代理模型"],
     ["2023/11", "代理模型只用當次資料更新", "loss 上升後來回震盪", "加入資料暫存"],
     ["2023/11", "圖形與響應存成暫存，每次用全部資料重訓", "loss 明顯下降並趨平", "定案，後兩代沿用"]],
    notes=[
        "出處：陳廷茂 2023/07–11 的週報、會議紀錄與版本紀錄。",
        "左圖：加入圖形與響應暫存後的 loss，2023/11/29 週報，28 GHz、25×25、只算 26.5–29.5 GHz，當時的 loss 定義和碩論不同。依原圖讀值，約 200 次內降到接近 0；只用當次資料的版本 loss 上升到約 230 後震盪。",
        "直接反傳：用 numpy 自寫全連接網路，換過多種激活函數、初始化、正規化與二值化方式；週報記錄神經網路無法順利反向傳播，之後轉 PyTorch，直通估計保留進碩論。",
        "強化學習：圖形當狀態、改像素當動作，獎勵設計改了 14 版；28 GHz 25×25 的獎勵前 600 步大幅起伏，之後停在固定值附近，依原圖讀值。之後又參考強化學習結合模仿學習的天線設計文獻，試 DDPG、Dueling DQN、PPO 加預訓練，4.7 GHz 小天線約跑 2,100 步。",
        "4.7 GHz 驗證題：16×20 格、每格 0.1 mm，規格 4.6–4.8 GHz、增益 1.50 dBi、反射 −10 dB，用確定能達標的題目先檢查算法。",
        "同期也向 Ansys 詢問加速方式，回覆是頻域求解的貼片天線用 GPU 幫助有限。",
    ])

# ================================================================ 9 第一代嘗試歷程：設計空間與饋入固定金屬
attempts_fig_slide(
    "第一代嘗試歷程：設計空間與饋入固定金屬",
    "從 4.7 GHz 驗證題到 28 GHz 雙層板、再到四層封裝板，網格大小與饋入固定金屬的位置都試過多種配置；左圖為四層板第一層的兩種固定區位置。",
    pair_patterns([A1 / "b1_4layer_m1_fixed_bottom_2024-02.png", A1 / "b2_4layer_m1_fixed_center_2024-02.png"],
                  ["固定區在下方", "固定區移到正中間"]),
    [["2023/10–11", "網格 16×20、20×20、50×50、25×25", "格子越細，模擬越久", "定 25×25、每格 0.2 mm"],
     ["2023/11–12", "饋入點周圍 8×8 格強制放金屬，並做有無對照", "不固定時 loss 震盪較大", "之後各版都固定"],
     ["2023/12", "四層板初版：第一層 3×3 mm、12×12 格", "有朝目標範圍逼近", "調整格數與區域"],
     ["2024/01–02", "第一層改 15×15，再擴成 5×5 mm、25×25", "與雙層板同大", "碩論採 25×25"],
     ["2024/01–02", "第一層固定區放大、下移、移到正中間", "loss 仍有大幅震盪", "碩論採中央偏下，對準 via"],
     ["2023/12 起", "第二層正對第一層處強制留空", "沿用既有四層板配置", "碩論沿用"]],
    img_w=4.9, trimmed=False,
    notes=[
        "左圖：2024/02/27 週報的四層板第一層，中間那塊整塊金屬就是饋入固定區；兩次都用上下界遮罩，依原圖讀值 loss 都還有大幅震盪。原圖未改，下方標籤是我們加的。",
        "饋入固定金屬的做法：出圖網路輸出後，再疊上一張固定區陣列。2023/11/19 在 4.7 GHz 版首次加入；2023/12 拿掉固定區重跑當對照，loss 震盪較大，成為碩論實驗一，固定後 loss 較快、較穩地下降。口試時委員問能否在生成時就給限制，而不是出圖後再覆寫。",
        "網格：28 GHz 比過 20×20、50×50 每格 0.1 mm 與 25×25 每格 0.2 mm，定為 25×25、每格 0.2 mm。",
        "四層板沿用林佑宣的四層封裝板配置：第一層初版 3×3 mm、12×12 格，第二層 12×5.5 mm、48×22 格、每格 0.25 mm，第二層正對第一層處留空。碩論採第一層 25×25、第二層 48×22，固定區在中央偏下、對準 via。",
        "對稱與形狀拼接第一代沒有做；選逐格的像素網格，是為了不被特定幾何形狀綁住。",
    ])

# ================================================================ 10 第一代嘗試歷程：loss、代理模型與方向圖條件
attempts_fig_slide(
    "第一代嘗試歷程：loss、代理模型與波束寬條件",
    "流程定案後，第一代持續調整 loss 寫法與代理模型結構，也試過加入波束寬條件；碩論以反射係數與增益為優化目標。左圖為加入波束寬條件的一次運行。",
    A1 / "c2_beamwidth_4layer_2024-04.png",
    [["2023/11–12", "loss 只罰未達標頻點，再加上下界遮罩", "會議紀錄記增益項改善", "碩論的兩種 loss"],
     ["2024/02–03", "代理模型每層都加 Batch Normalization", "梯度變小，loss 停在固定值", "改成只加後兩層"],
     ["2024/03", "代理模型的 loss 改用 MAE", "對 loss 下降影響不明顯", "維持 MSE"],
     ["2023/11–2024/05", "激活函數 Tanh、LeakyReLU、PReLU 比較", "PReLU 收斂最快", "定 PReLU"],
     ["2024/04–05", "加波束寬條件：兩切面 ±30° 內高於 1 dBi", "加條件後 loss 仍能下降", "場型放附錄，列為未來方向"],
     ["2024/07", "口試後：出圖網路改卷積、做連通分析", "對角相接的像素實際可能斷開", "交接給第二代"]],
    img_w=4.6,
    notes=[
        "左圖：2024/04/23 週報，四層板加波束寬條件的一次運行；上排是 Phi 0° 與 90° 兩切面場型，紅虛線是 ±30°、藍點線是 1 dBi，下排是 S11 與增益。原圖黑框換成白底，內容未改。",
        "波束寬條件：θ 在 −30° 到 30° 增益要高於 1 dBi，週報寫是實驗室學長要求同時看兩個切面。當年碩論以反射係數與增益為優化目標，方向圖不在目標內，定義不同，不構成評價；碩論結論把增加輻射場型條件列為未來方向。",
        "loss：原本用 MSE 逼近梯形目標曲線，已達標的頻點也被罰；改成只罰未達標頻點，再加 24–32 GHz 上下界遮罩，增益目標由 5 改為 4 dBi，成為碩論的兩種 loss。",
        "代理模型的其他嘗試：兩層輸入串接或分支，用 200 筆已存資料離線比較，訓練 loss 0.26 對 0.34，碩論採串接；每次迭代改為從頭重訓；接續前次模型與資料再訓。",
        "口試後：委員建議代理模型考慮 CNN、補驗證集 loss，並指出對角相接的像素在模擬中算連通、實際製作可能斷開。學長把從饋入點連出的金屬塗色，若對角不算連通，有效金屬明顯變少；並提出內部格子四邊各擴 0.05 mm、讓對角重疊的構想，交接給第二代。",
    ])

# ================================================================ 11 第一代成果
fig_table_slide(
    "第一代：四組模擬設計與規格對照",
    "雙層板與四層封裝板各做 n257 天線與濾波天線；n257 兩組只差一到三個頻點，濾波兩組帶外仍有較多頻點未達。左圖為四層板濾波天線。",
    T1 / "tab4-5_four_layer_filter_result.png",
    [["設計", "迭代次數", "未達處，依原圖判讀"],
     ["雙層板 n257", "1,000", "29.5 GHz 增益略低於 4 dBi"],
     ["四層板 n257", "1,500", "26.5 GHz S11、29–29.5 GHz 增益略差於門檻"],
     ["四層板濾波", "1,500", "帶內增益 3 點、帶外 S11 2 點、帶外增益 3 點"],
     ["雙層板濾波", "約 1,000", "帶內增益 4 點、帶外 S11 6 點、帶外增益 2 點"],
     [("留給後代", True), "25×25 像素化設計、生成網路加代理模型的線上學習、饋入固定金屬、只罰未達標頻點的 loss", None]],
    [1.55, 1.0, 5.22], img_w=4.1, row_h=0.72, body_size=13.5, center_cols=(1,),
    notes=[
        "出處：陳廷茂碩論表 4-3～4-6。未達處是我們依碩論原圖讀值，誤差約 ±0.30 dB：雙層板 n257 在 29.5 GHz 增益約 3.50 dBi；四層板 n257 在 26.5 GHz S11 約 −9.30 dB、29–29.5 GHz 增益約 3.80–3.90 dBi。",
        "左圖是四層板濾波天線：第一層 25×25、第二層 48×22 格可調，饋入點周圍 8×8 格固定放金屬；紅線上界、淺藍下界、黃底 n257 頻段。",
        "規格：n257 天線 26.5–29.5 GHz S11 < −10 dB、增益 > 4 dBi；濾波天線另在 24–32 GHz 要求帶外高反射、低增益，訓練時以上下界表示。",
        "每換一組規格要重跑一次完整迭代。全部是 HFSS 模擬，當年沒有打樣，方向圖未列入優化。",
    ])

# ================================================================ 第二代：12 方法骨架
fig_text_slide(
    "第二代：兩項方法",
    "BiScaleNorm 讓二值化訓練不再卡住；二階段生成以週期翻轉像素跳出局部解。左圖為第二階段的結果。",
    T2 / "fig4-8_stage2_result.png",
    [("BiScaleNorm：穩定二值化訓練", ["金屬圖形只能是 0 或 1，訓練容易停住、依賴學習率",
                                  "正負兩側分開縮放後，三種學習率都能收斂"]),
     ("二階段生成：跳出局部解", ["第一階段先收斂到一個局部解，弱點在 29 GHz", "第二階段每 400 次翻轉少量像素，沒改善就退回歷史最佳圖",
                          "29 GHz 的反射係數與增益都比第一階段改善"])],
    img_w=6.4, size=16,
    notes=[
        "出處：錢鵬予碩論〈基於二階段梯度探索之貼片天線金屬圖形生成方法〉§4.2.1、§4.2.2，圖 4-4～4-8，2025/06。",
        "BiScaleNorm：學習率 0.003 與 0.0003 時原設計 loss 停住、BiScaleNorm 穩定收斂；0.00003 時兩者都收斂，依碩論判讀 BiScaleNorm 約第 400 次、原設計約第 1,000 次。",
        "左圖：(a) 金屬圖形、(b) 第二階段 loss、(c) 反射係數、(d) 增益，黑線為規格上下界、紅線為模擬結果。變異率起始 5%、週期 400，都是經驗值。",
        "全部是 HFSS 模擬；loss 為該論文自訂，不和其他世代比較。接下來三頁是第二代走到這兩項方法的嘗試。",
    ])

# ================================================================ 13 第二代嘗試歷程：時間軸
fig_slide(
    "第二代嘗試歷程：一年的時間軸",
    "2024 年 8 月接手第一代流程後，每月嘗試一到三個方向，2025 年春收斂為 BiScaleNorm 與二階段生成；圖為學長自繪的時間軸。",
    A2 / "ga2_timeline_his.png", pad=0.3,
    notes=[
        "圖為錢鵬予學長自己整理的一年時間軸，2025/06 研究彙整第 2 張，原圖。",
        "圖上 GEN 是代理模型；Custom Normalize 即後來的 BiScaleNorm；Mutation 是週期翻轉像素的變異；Selection Mode 是每格從幾種小圖案中挑一個的拼接方式；Sign 函數探討是近似二值化；Random Sampling 是生成器改用隨機輸入；Opening & Closing 是處理孤島的開閉運算；Gradient Doping 是在梯度上加雜訊的構想。",
        "很多新方法先在 RIS 電磁面鏡的相位分布任務上試：RIS 的響應用公式直接算，不必跑 HFSS、也不需要代理模型，一輪約 10–20 小時，適合當試驗台；RIS 任務上可行，不代表接上代理模型的貼片任務也可行。",
        "2024/12 實驗室的 HFSS 整月無法使用，該月以 RIS 任務與文獻為主。",
    ])

# ================================================================ 14 第二代嘗試歷程：生成器與訓練方法
attempts_slide(
    "第二代嘗試歷程：生成器與訓練方法",
    "針對梯度消失與局部解，第二代先試換架構、二值化與優化器，最後以正規化和翻轉像素兩條路收斂。",
    [["2024/08–12", "生成器改用 CNN、LSTM", "CNN 出現棋盤紋或大片色塊；LSTM loss 停住", "維持全連接生成器"],
     ["2024/09–12", "以近似函數取代硬性二值化", "較穩，但送模擬前仍要硬性二值化", "改從正規化著手"],
     ["2024/10–2025/01", "比較六種優化器；loss 改成只罰超出上下界", "RMSProp 優於 Adam，SGD 發散", "上下界加 Smooth L1 成為碩論 loss"],
     ["2024/12–2025/03", "非對稱正規化，再加極值平滑", "重複多次、不同學習率都能下降", "BiScaleNorm，碩論方法一"],
     ["2024/09–2025/02", "權重平均、熱重啟、隨機輸入、判別器", "圖形骨架前幾百次就定型，調參數難以改變", "翻轉像素加回滾，碩論方法二"],
     ["2025/02–06", "檢查代理模型的擬合與梯度", "新資料擬合變差，導數有誤差", "列為未來方向"]],
    col_w=(1.75, 3.75, 3.6, 3.07),
    notes=[
        "CNN 與 LSTM：CNN 生成器試過反卷積、上採樣、先全連接再卷積三版，在 RIS 任務上分別出現棋盤紋與大片色塊；貼片任務另比過有無 CNN。4 層 LSTM 加 1 層全連接時 loss 幾乎不下降。",
        "近似二值化：ApproxSign、Swish-like、Smooth Thresholding；RIS 任務上 ApproxSign 最穩，學長也用過一段時間，因送模擬前仍要再硬性二值化，改從正規化著手。",
        "優化器：比較 Adam、AdaBelief、NAG、RMSProp、AdaBound 與 SGD；碩論最後用 Adam，並以文獻說明二值化空間中 Adam 優於 SGD。loss 先由 MSE 改 Smooth L1，再改成上下界遮罩。",
        "正規化：正值除以最大值、負值除以最小值的絕對值，把數值壓回有梯度的範圍；RIS 任務重複多次都穩定下降，對照的 LayerNorm 多次停住。2025/03 極值改用指數移動平均，命名 BiScaleNorm。",
        "局部解：權重平均觀察到會跳出又回到相同位置；隨機輸入 loss 最低但迭代次數很高，與變異不相容；判別器與好壞分類器看不出明顯幫助。學長觀察到圖形骨架在前幾百次迭代就定型、之後只改細節，這是翻轉像素的動機。",
        "代理模型：越晚加入的資料擬合越差，1000 筆後重訓要 60000 次迭代；小測中經代理模型擬合比直接擬合慢且抖動。學長建議把梯度誤差當指標，代理模型改為少量事先蒐集加線上蒐集；第四代從前代累積的資料起步，是同一個方向。",
        "另有只出現在程式裡的變體，例如 DyT、Ranger、殘差生成器，本報告不展開。",
    ])

# ================================================================ 15 第二代嘗試歷程：圖形約束與新結構
attempts_fig_slide(
    "第二代嘗試歷程：圖形約束與新結構",
    "對稱、小圖案拼接、饋入固定金屬、孤島與對角導通都試過，留下饋入固定金屬、保留孤島、0.21 mm 格寬與雙埠濾波結構；左圖為刪除孤島前後的響應。",
    A2 / "ga2_island_delete_compare.png",
    [["對稱", "2024/09", "以對稱方式生成，想改善場型凹陷", "第三、四代接續"],
     ["小圖案拼接", "2024/09、2025/05", "每格從 6–8 種預設小圖案挑一個", "第三代七巧板延續"],
     ["饋入固定金屬", "全程", "每次生成後在饋入端疊一塊固定金屬", "寫入碩論，四代都保留"],
     ["孤島", "2025/02–04", "形態學處理、刪除前後比較", "刪除會改變響應，決定保留"],
     ["對角導通", "2025/04", "模擬時每格畫成 0.21 mm，對角重疊導通", "0.21 mm 成為後代模擬格寬"],
     ["結構與頻段", "2025/05–06", "雙埠濾波結構、換到 n260", "雙埠成為後代濾波主線"]],
    head=["主題", "時間", "做法", "後續"], split=(0.2, 0.19, 0.35, 0.26), img_w=4.6,
    notes=[
        "左圖：學長口委答覆的對照，上排保留孤立金屬、下排手動刪除，由左至右是金屬圖形、反射係數、增益。",
        "對稱：2024/09 以對稱方式生成，目的是提高可解釋度、改善場型凹陷，留下 S11、增益曲線與 loss；之後的簡報轉向正規化與變異，第三、四代接續對稱。",
        "小圖案拼接：2024/09 版每格從 6 種 3×3 小圖案挑一個；2025/05 重啟，4×4 與 2×2 兩組各 8 種。兩次都先在 RIS 任務上試，重啟版在 RIS 上響應貼近目標；接上代理模型的貼片任務 loss 沒有持續下降，學長列為處理好代理模型後的方向。第三代的七巧板延續這個想法。",
        "饋入固定金屬：沿用第一代，每次生成後在饋線接點疊一塊固定金屬；碩論圖上 7×7 格，程式是 8×7；雙埠改上下兩塊。",
        "孤島：8 鄰接標記獨立金屬群，試過膨脹、侵蝕、開運算、閉運算，也試過把變異換成開閉運算；不同處理送 HFSS 響應都會變。2025/04 比較刪除前後，孤島多或靠近饋入時影響更大，決定保留；之後的報告寫不刪孤島後，每次訓練都能得到大致符合的結果。",
        "對角導通：填金屬時每格由 0.2 mm 放大到 0.21 mm，對角相接的像素重疊而導通，寫入碩論附錄二；合作實驗室另提六角格，當時列為低優先。",
        "雙埠與 n260：微帶線從圖形兩側穿過、各 7.5 mm，看 S11、S22、S21，寫入碩論附錄一，成為第三、四代濾波器的主線；n260 只改 HFSS 模擬設定，學長記錄結果與趨勢和 n257 一致。雙埠結構對應 2025/05/01 會議中討論的從架構上改動濾波天線。",
    ])

# ================================================================ 第三代：16 方法骨架
fig_text_slide(
    "第三代：單階段線上學習與三項機制",
    "沿用生成器加代理模型的線上學習，加入三項機制，並以理論公式算出的 28 GHz 貼片當起點。",
    T3 / "fig3-11_design_flow.png",
    [("三項機制", ["ACP：訓練停滯時自動重新拉高學習率，把前代的兩階段合成一段",
                "DLF：只挑表現優於累計平均的樣本給代理模型複習",
                "SC loss：訓練中要求金屬連成一塊，抑制孤島"]),
     ("起點", ["用傳輸線公式算出 28 GHz 貼片當初始圖形，不從空白開始"])],
    img_w=7.0, size=16,
    notes=[
        "出處：吳維文碩論〈自適應循環策略與圖譜連通度損失函數於貼片天線金屬圖形生成之研究〉圖 3-11，2026/06。",
        "規格：n257 26.5–29.5 GHz，S11 < −10 dB、單一方向增益 > 4 dB。",
        "每次迭代仍跑 HFSS；代理模型提供梯度，不取代模擬。接下來四頁是第三代走到這三項機制的嘗試。",
    ])

# ================================================================ 17 第三代嘗試歷程：七巧板與對稱
attempts_fig_slide(
    "第三代嘗試歷程：七巧板與對稱",
    "2025 下半年在逐格生成之外試了兩種生成方式：從形狀庫挑小形狀拼滿，以及把生成圖做對稱翻轉再挑最好的候選；左圖為 35 種形狀庫。",
    T3 / "ase_202511_tangram_shape_library.png",
    [["2025/07–10", "七巧板：8 種粗形狀，再加細線、十字與對稱形狀", "推測可選形狀太少，逐步增加", "擴成形狀庫"],
     ["2025/11–12", "七巧板：35 種形狀庫，比較兩種可微挑法", "Sinkhorn 對排列與選擇的約束較嚴", "2026 起改以逐格生成為主"],
     ["2025/07", "對稱：左右翻轉，擴增代理模型的資料", "兩張翻轉圖之一較好", "擴成兩軸"],
     ["2025/10–11", "對稱：兩軸 8 種翻轉候選", "全送模擬成本高", "改由代理模型先篩"],
     ["2025/12", "對稱：生成模型直接產生再鏡射", "能生成對稱結構", "2026 起轉向孤島抑制"]],
    img_w=5.0, row_h=0.78,
    notes=[
        "左圖：2025/11 月報的 35 種形狀庫，含粗細線、十字、X、點、圓與 T、L、U 型各四個旋轉。",
        "七巧板：從形狀庫挑 5×5 小形狀拼滿 25×25，原意是用粗形狀避免金屬不導通，想法源自第二代的小圖案拼接。第一版 8 個粗形狀，學長推測可選形狀太少；2025/10 依會議建議加細線、斜線、十字、棋盤、螺旋到 14、16 種，並另做 11 種上下左右都對稱的形狀；2025/11 擴成 35 種，比較 Gumbel-Softmax 與 Gumbel-Sinkhorn 兩種可微挑法。",
        "2025/12 月報用七巧板與逐格兩種模式比較兩種重啟方式；2026/01 起月報以逐格模式為主，2026/05 的程式仍保留七巧板的訓練檔。",
        "對稱：2025/07 先把雙埠濾波器的生成圖左右翻轉，一起送模擬訓練代理模型，目的是讓代理模型看到更多樣的圖形。2025/10/02 的報告記有天線至少一軸對稱的討論；同月擴成兩軸 8 種翻轉，起初 8 個候選全送 HFSS，每輪多 8 次模擬；10/30 改成代理模型先預測、只模擬最好的 1 到 2 個；11/27 改成最好與最差都模擬。月報上 8 個候選的 loss 是代理模型預測值。",
        "2025/12 換成條件變分自編碼器直接產生圖形、再鏡射成 6 個候選，低權重孤島抑制時能生成對稱結構。口委問答補充：場型不對稱的根本成因是整體幾何不對稱。",
    ])

# ================================================================ 18 第三代嘗試歷程：從兩階段到單階段訓練
attempts_fig_slide(
    "第三代嘗試歷程：從兩階段到單階段訓練",
    "把前代先找局部解、再週期變異的兩階段流程，逐步改成一段訓練內自動判斷停滯並重啟，同時調整 loss、代理模型的複習資料與起點；左圖為 ACP 的一個週期。",
    FIG / "gen3" / "g3_acp_lr_tau_cycle.png",
    [["2025/07", "變異改成停滯 200 次自動觸發", "一次收斂即可進入變異", "改用溫度控制探索"],
     ["2025/10", "loss 只罰不達標的頻點", "目標頻段更貼近規格", "碩論的邊界 loss"],
     ["2025/10–2026/01", "停滯就重啟，比較三種重啟方式", "linear 較好", "碩論 ACP"],
     ["2025/10–2026/02", "代理模型只複習優於平均的樣本", "同為 1000 次，最小 loss 較低", "碩論 DLF"],
     ["2025/11", "理論公式算出的 28 GHz 貼片當起點", "起點已在 28 GHz 諧振", "碩論採 Base-2"],
     ["碩論", "生成器試 CNN、條件變分自編碼器、GAN", "小樣本下以多層感知機較穩", "沿用多層感知機"]],
    img_w=4.7,
    notes=[
        "左圖：ACP 的一個標準週期，學習率與二值化溫度同步，先線性暖身、再餘弦退火；高溫時圖形邊界模糊、方便探索，低溫時銳化成 0 與 1。碩論設定週期 100 步、暖身占兩成、停滯判定 25 步。",
        "變異自動觸發：連續 200 次 loss 沒有進步就變異，鄭國宏學長原本建議 100 次；學長認為週期變異偏向運氣，2025/10 起改成偵測停滯就重啟的循環排程，當時叫 ACS，碩論改稱 ACP。五種排程各跑 500 次比較，ACP 最小 loss 2.07 最低，碩論表 4-2；loss 是第三代自己的定義，0 代表帶內全部達標，不跨代比較。",
        "重啟方式比較 reset、peak、linear，逐格模式下 linear 較好，碩論採 linear。",
        "代理模型的複習資料：先試只回放優於平均的樣本、全資料定期重訓、對稱候選最好與最差都送模擬；2026/02 定稿為 DLF，所有樣本都存，每次重訓前用歷史 loss 的累計平均重新篩，同為 1000 次訓練，最小 loss 由約 3.60 降到 1.71。",
        "熱啟動：用傳輸線公式與 RO4003C 參數算出 28 GHz 貼片，取整到 0.2 mm 像素時捨去得 Base-1、進位得 Base-2，碩論採 Base-2；公式起點由鄭國宏學長提供，Base-2 在 28 GHz 的 S11 約 −25 dB、增益約 5 dB。",
        "生成器：碩論寫試過 CNN、條件變分自編碼器與 GAN，在模擬成本高、樣本少的條件下，選擇沿用多層感知機。",
    ])

# ================================================================ 19 第三代嘗試歷程：孤島與連通
attempts_fig_slide(
    "第三代嘗試歷程：孤島與連通",
    "為了減少不連到饋入點的金屬，先後試了四種可微的孤島懲罰與兩種事後修補，並定義饋入連通率來量化；左圖綠色為與饋入點相連的金屬。",
    A3 / "g3a_island_loss_tv_gc_is.png",
    [["2025/12–2026/01", "TV：懲罰相鄰像素差異", "只改善饋入點附近", "碩論列為局部法"],
     ["2026/01", "IS：比較每格與周圍平均", "與 TV 相近", "碩論改稱 KNS"],
     ["2026/02", "GC：把閉運算寫進 loss", "連通有改善", "由 SC 取代"],
     ["2026/02–05", "SC：以整張圖的連通度當 loss", "連通能延伸到上半部", "碩論採用"],
     ["2026/02–04", "事後修補：閉運算、少量補洞", "響應或場型會變", "改在訓練中約束"],
     ["2026/03", "定義饋入連通率", "不可微，只用來評估", "碩論沿用"]],
    img_w=5.0,
    notes=[
        "左圖：同一張圖互比的饋入連通率，綠色是從饋入點只走上下左右能到的金屬，灰色是孤島，紅點是饋入點；Base-2 13.75%、TV 權重 50 與 100 為 18.53% 與 59.62%、GC 100 為 53.95%、IS 權重 1 與 100 為 33.02% 與 61.00%。",
        "TV 懲罰相鄰像素差異，2025/12 首次加入時學長記錄 S11 與增益保持穩定、圖案明顯更連續；IS 比較每格與周圍一圈的平均，碩論改稱 KNS；學長整理 TV 在天線較好、IS 在濾波較好。",
        "GC：先把膨脹再侵蝕的閉運算當後處理，發現會刪到對響應有貢獻的獨立金屬，才把閉運算寫進 loss。SC：把像素當成圖，取拉普拉斯矩陣第二小特徵值的倒數當 loss，整張圖一起算；低、高權重的饋入連通率 31.13%、87.80%，權重極大時 98.18% 但 S11 與增益不達規格，碩論取平衡權重。",
        "事後修補：少量補洞把大孤島接上，濾波器只能補少量、特性才不會大改，單埠補了以後場型更不對稱，所以改在訓練中用 loss 約束。碎圖形還會讓 HFSS 網格劃分出錯，有抑制孤島時幾乎不發生。",
        "饋入連通率只算上下左右相連，對角不算，理由是對角接點電流貢獻小、蝕刻時容易斷；這和第四代處理的對角導通是同一件事。",
    ])

# ================================================================ 20 第三代嘗試歷程：饋入固定區與濾波器
attempts_fig_slide(
    "第三代嘗試歷程：饋入固定區與濾波器",
    "饋入點旁的固定金屬區比較過有無與大小；雙埠濾波器在暫停後恢復，並以 loss 讓兩個饋入點接通；左圖綠色為與饋入點相連的金屬。",
    pair_patterns([Image.open(T3 / "ase_ccu0528_feed9x9_rfeed52.png").crop((100, 0, 720, 582)),
                   Image.open(A3 / "g3a_filter_loss_rfeed4.png").crop((640, 0, 1360, 535))],
                  ["單埠，9×9 固定區", "濾波器，TV 加區間 loss"], label_px=46, gap=50),
    [["2025/07–09", "濾波器比較有無固定饋入金屬", "當時不固定的那組較好", "之後一律保留"],
     ["2026/05", "單埠固定區 5×5 擴為 9×9", "饋入連通率 52.42%", "碩論版預設仍 5×5"],
     ["2025/10–2026/01", "先做天線，濾波暫停後恢復", "TV、IS 各有所長", "改用區間 loss"],
     ["2026/03", "濾波器加 TV 與區間 loss", "兩個饋入點接通", "碩論改驗證泛用性"],
     ["2026/06", "ACP、熱啟動、DLF 原樣移到濾波器", "碩論寫初步符合規格", "碩論實驗四"]],
    img_w=5.0, row_h=0.78, trimmed=False,
    notes=[
        "左圖由兩張學長原圖各取一格並排，下方標籤是我們加的。左格出自 2026/05/28 月報，單埠固定區擴為 9×9 時的連通區，饋入連通率 52.42%；右格出自 2026/03 的 AI 模型設定，雙埠濾波器 TV 100 加區間 loss，紅點是上下兩個饋入點，饋入連通率 94.03%，同一張圖裡 SC 200 與舊 loss 的 TV 100、IS 100 都是 0.00%。都是學長報告值。",
        "有無固定饋入金屬：2025/07–09 雙埠變異實驗中，學長備註不固定的那組較好；當時生成器輸出本身沒有套固定區，只在翻轉或變異後才套上，條件和後來不同。之後一律保留固定區，單埠在下緣中央、雙埠上下各一塊，程式寫明都是 5×5。",
        "9×9：2026/05 為了改善連通把單埠固定區放大到 9×9，饋入連通率 52.42%、S11 與增益初步達標，上方仍有孤島；同月程式快照的預設仍是 5×5。",
        "濾波器：2025/10/02 報告列出先做天線、暫停濾波，2026/01 恢復，比較 TV 1、TV 100、IS 100；2026/03 起比較舊 loss 與區間 loss，TV 加區間 loss 讓兩個饋入點接通。",
        "碩論實驗四把 ACP、熱啟動、DLF 原樣移到雙埠濾波器，超參數不重調、不加孤島懲罰，用來驗證方法換規格也能用，碩論寫初步符合電磁規格。",
    ])

# ================================================================ 21 第三代成果：逐項比較
fig_table_slide(
    "第三代：逐項加入的比較與孤島抑制",
    "逐項加入三項機制，各列條件不同、只作定性參考；響應 loss 為 0 代表帶內全部達標，加入 SC loss 後饋入連通率由 14.01% 提高到 61.98%。",
    T3 / "tab4-3_base_tv_sc_radiation.png",
    [["設定", "響應 loss", "饋入連通率"],
     ["固定學習率，不加機制", "3.64", "27.22%"],
     ["只加 DLF", "1.71", "16.28%"],
     ["只加 ACP", "2.07", "18.89%"],
     ["ACP 與 DLF，左圖 Base", "0.00", "14.01%"],
     [("三項合用，左圖 SC", True), ("0.99", True), ("61.98%", True)]],
    [2.95, 1.55, 1.75], img_w=5.6, row_h=0.62, body_size=15, center_cols=(1, 2),
    notes=[
        "出處：吳維文碩論表 4-4、表 4-3、圖 4-8、圖 4-9。",
        "饋入連通率：從饋入點只走上下左右能連到的金屬占全部金屬的比率，對角相接不算連通。左圖綠色是連通的金屬、灰色是孤島，下排是 28 GHz 場型。",
        "響應 loss 只算超出規格的部分，是第三代自己的定義，不跨代比較。只加 ACP 的一列訓練 500 次、只加 DLF 的一列 1,000 次。",
        "左圖三種設定都用理論貼片起點、ACP 與 DLF，差別只在孤島抑制：不加、TV loss 22.39%、SC loss 61.98%；SC 權重過大時連通率可到 98.18%，但 S11 與增益不達標。",
    ])

# ================================================================ 22 第三代成果：樣品量測
fig_text_slide(
    "第三代期間的 5 個實作樣品量測",
    "2026/03/24 聯合審查報告 5 個單元天線樣品的無反射室量測；量測通帶比模擬整體往高頻移。",
    T3 / "ase_202603final_sample1_s11.png",
    [("回波損耗 > 10 dB 的通帶，模擬 → 量測", [
        "樣品 1：26.1–30.0 → 26.9–31.6 GHz", "樣品 2：26.2–30.1 → 26.9–31.6 GHz",
        "樣品 3：26.2–29.9 → 26.7–31.4 GHz", "樣品 4：26.1–30.2 → 26.7–31.6 GHz",
        "樣品 5：26.2–30.1 → 26.7–31.5 GHz"]),
     ("觀察", ["量測下緣 26.7–26.9 GHz，未涵蓋 n257 下緣 26.5 GHz"]),
     ("留給後代", ["饋入連通率指標、連通 loss、理論貼片起點，以及四代唯一的實測經驗"])],
    img_w=6.2, size=15,
    notes=[
        "出處：2026/03/24 ASE-CCU 聯合審查報告。左圖是樣品 1 的 S11，紅虛線模擬、藍線量測；樣品 1 的量測在更早一次聯合審查已先報告過。",
        "模擬已含 K 接頭與 ENIG 化金層；量測在無反射室進行。由報告數字推算，量測下緣高 0.5–0.8 GHz、上緣高 1.4–1.6 GHz。",
        "28 GHz 峰值增益模擬／量測：樣品 1 至 5 為 2.20／3.00、2.00／2.70、1.30／2.70、0.90／2.10、1.50／2.40 dBi，量測都高於含化金層的模擬。",
        "2026/05/28 月報：顯微鏡量得實作像素 0.188–0.231 mm，模擬為 0.21 mm。樣品的設計流程、設計者與過蝕刻分析本報告未納入，待確認。",
    ])

# ================================================================ 第四代：23 接手的缺口
columns_slide(
    "第四代：接手後接續的三件事",
    "前代留下生成器、代理模型與 34,215 筆模擬，第四代各任務都從前代累積的設計出發；當年目標不含方向圖與可製造性，我們接手後接續這三件事。",
    [("可製造", ["懸浮金屬與 1–3 像素的小碎片", "對角相接的像素只靠 0.01 mm 重疊相連", "依一般 PCB 規則有導通疑慮，請日月光確認",
               "直接刪碎片，響應會大幅變差"]),
     ("規格擴充", ["方向圖：28 GHz 兩切面 ±45° 內不低於 0° 方向 −3 dB", "金屬左右對稱的天線", "雙埠濾波器與 11–40 GHz 寬頻濾波器"]),
     ("效率", ["單機逐筆：生成、模擬、訓練", "每筆模擬後等代理模型訓練", "調整算法時模擬機停下來等"])],
    size=17,
    notes=[
        "前代資料：承接的歷史模擬 34,215 筆，單埠 24,189、雙埠約 1.0 萬。單埠天線與雙埠濾波器的起點直接取自前代資料；對稱天線與寬頻濾波器的起點是由前代設計演變而來的我們較早的設計。",
        "小碎片與懸浮金屬的處理結果見本段最後的可製造頁。",
        "對角接點：模擬時每格金屬畫成 0.21 mm，對角相鄰處只重疊 0.01 mm。",
        "當年的優化目標不含方向圖與可製造性，定義不同，不構成評價。",
    ])

# ================================================================ 24 第四代嘗試歷程：從逐筆到批次線上學習
attempts_slide(
    "第四代嘗試歷程：從逐筆到批次線上學習",
    "6–7 月先復現前代流程並調整代理模型，進步仍靠偶發跳躍；之後改從前代設計出發，並把模擬與訓練分開。",
    [["6 月上旬", "復現前代流程，加可學習輸入、邊界損失、對稱生成器", "差異小於線上學習本身的起伏", "三項停用"],
     ["6–7 月", "方向圖納入評分", "帶內 +0.44 dB 的前代設計，方向圖還差 2.71 dB", "方向圖成為獨立條件"],
     ["6/26–7/10", "五種代理模型更新方式，移除回滾", "最佳 −2.89 dB，來自一次跳躍", "逐筆線上學習收束"],
     ["7/03–7/06", "從前代資料出發", "前代資料池已有通過規格的設計", "各任務從前代設計出發"],
     ["7/03 起", "批次線上學習", "一台一天由約 37 筆變約 540 筆", "第四代主流程"],
     ["7/06 起", "同一設計多次、多台機器重跑", "同設定重跑結果完全一致", "紀錄須多台機器重跑一致"]],
    notes=[
        "6 月三項是 6/24 給日月光六月報告裡的探索方向；loss 是前代框架的複合 loss，不是 dB。前代 41 條軌跡跑到 500 輪，同一方法跨軌跡的標準差約 1.50 dB，所以小改動比不出好壞。可學習輸入的梯度傳回輸入端幾乎消失；對稱後來改成直接作用在圖形上。",
        "方向圖要求：28 GHz 兩切面 ±45° 內不低於 0° 方向增益 −3 dB，6/24 由 ±55° 改為 ±45°，7/07 起成為與帶內並列的獨立條件。當年目標不含方向圖，定義不同，不構成評價。",
        "五種代理模型更新方式：只訓精選樣本一輪、訓到收斂、全部資料重訓、多模型加信任門控、自適應與滑動視窗訓練量。−2.89 dB 只算 S11 與增益，來自一次 +4.24 dB 的跳躍；訓練誤差壓下來後最佳仍停在 −3.65 dB，判斷瓶頸在候選從哪裡來。移除回滾的理由：退回最佳狀態的規則讓搜尋卡在第一個山頭；退回舊生成器又配上持續更新的代理模型，前後不一致。",
        "前代資料：零模擬估算顯示逐筆優化到規格約需 2×10⁵ 輪，前代資料池內已有 +0.38 dB 的通過設計；18 個前代通過設計以目前的模擬設定重跑，8 個仍通過，最佳 +0.44 dB，模擬設定與當年不同。單埠天線與雙埠濾波器的起點直接取自前代；對稱天線與寬頻濾波器的起點是由前代設計演變而來的我們較早的設計。",
        "批次線上學習：單埠天線每筆中位約 160 秒，一台一天約 540 筆；逐筆線上學習時每輪約 40 分鐘，一台一天約 37 筆。",
        "重跑：同設定重跑結果位元級一致，41 次全同值。7/06 一筆 +0.48 dB 的首個全過，重跑 8 次都是 −0.06 dB，原因是我們的後處理把增益曲線對錯半格，已改成按頻率對位。",
    ])

# ================================================================ 25 第四代：方法流程
s, top = content_slide(
    "第四代：批次線上學習流程",
    "代理模型負責大量初篩，人決定搜尋方向；研究者判讀與調整時，模擬機照常運轉。",
    notes=[
        "上圖是前代流程示意：產生一個候選、HFSS 模擬、訓練代理模型依序循環，每筆模擬後要等代理模型訓練完，研究者調整算法時模擬機也要等。",
        "下圖是批次線上學習：開發機一次挑出一組候選放進佇列，三台模擬機各自不停取件；收齊後重訓代理模型再挑下一組。所有數字以 HFSS 為準。",
        "速度：單埠天線每筆中位約 160 秒，一台一天約 540 筆，三台約 1,600 筆，是換算值。寬頻濾波器頻段寬，每筆中位約 464 秒，三台合計每小時約 22–24 筆、一天約 550 筆。到 10 月自行累積的 HFSS 結果超過 7.6 萬筆，含重複模擬。",
        "10/21 起寬頻濾波器在最好的一群設計之間改用實測統計排序，不再每輪用神經網路，見新版代理模型頁。",
        "每一輪開始前先寫下成功與否的判準。示意圖不按實際時間比例。",
    ])
lg_y = top + 0.05
lx = 6.25
for col, name in ((TEAL, "挑選或產生候選"), (NAVY, "HFSS 模擬"), (ORANGE, "訓練代理模型"), (GRAY, "人工調整")):
    rect(s, lx, lg_y + 0.09, 0.26, 0.2, col)
    textbox(s, lx + 0.3, lg_y, 1.75, 0.38, [name], size=13, color=INK)
    lx += 0.3 + _text_width(name, 13) + 0.28
y1 = top + 0.5
textbox(s, L - 0.04, y1, 11, 0.4, [("前代：線上學習，1 台模擬機", True, BLUE)], size=20)
x0, lab_w = L + 1.45, 1.4
textbox(s, L, y1 + 0.55, lab_w, 0.5, ["模擬機"], size=15, anchor=MSO_ANCHOR.MIDDLE)
bx = x0
seq = [(TEAL, 0.18), (NAVY, 0.5), (ORANGE, 0.34)]
for k in range(4):
    for col, w in seq:
        rect(s, bx, y1 + 0.55, w, 0.5, col)
        bx += w
gw = (R - x0) - 8 * 1.02
rect(s, bx, y1 + 0.55, gw, 0.5, GRAY)
textbox(s, bx, y1 + 0.55, gw, 0.5, ["人工調整算法，模擬機停下等待"], size=12.5, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
bx += gw
while bx + 1.02 <= R + 0.001:
    for col, w in seq:
        rect(s, bx, y1 + 0.55, w, 0.5, col)
        bx += w
y2 = y1 + 1.4
textbox(s, L - 0.04, y2, 11.8, 0.4, [("第四代：批次線上學習，3 台模擬機幾乎不空轉，單埠天線一天約 1,600 筆", True, BLUE)], size=20)
rows_y = [y2 + 0.55 + k * 0.6 for k in range(4)]
for name, yy in zip(("開發機", "HFSS 1", "HFSS 2", "HFSS 3"), rows_y):
    textbox(s, L, yy, lab_w, 0.44, [name], size=15, anchor=MSO_ANCHOR.MIDDLE)
span = R - x0
dev = [(0.0, 0.62, TEAL, "挑選"), (0.95, 1.2, ORANGE, "重訓代理模型"), (2.45, 2.0, GRAY, "人工判讀與調整"),
       (4.75, 0.62, TEAL, "挑選"), (5.7, 1.2, ORANGE, "重訓代理模型"), (7.2, 2.0, GRAY, "人工判讀與調整"),
       (9.5, 0.62, TEAL, "挑選")]
for off, w, col, txt in dev:
    w = min(w, span - off)
    rect(s, x0 + off, rows_y[0], w, 0.44, col)
    textbox(s, x0 + off, rows_y[0], w, 0.44, [txt], size=12.5, color=INK if col == GRAY else WHITE,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for k, start in enumerate((0.15, 0.32, 0.49)):
    bx = x0 + start
    while bx < R - 0.05:
        w = min(0.56, R - bx)
        rect(s, bx, rows_y[k + 1], w, 0.44, NAVY)
        bx += 0.62
ya = rows_y[-1] + 0.62
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(ya), Inches(R), Inches(ya))
ln.line.color.rgb = MUTED
ln.line.width = Pt(1.25)
lnel = ln.line._get_or_add_ln()
lnel.append(lnel.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
textbox(s, R - 1.0, ya + 0.02, 1.0, 0.35, ["時間"], size=13, color=MUTED, align=PP_ALIGN.RIGHT)

# ================================================================ 26 第四代：目前的系統架構
fig_slide(
    "第四代：目前的系統架構",
    "開發機產生並篩選候選，共享儲存排隊與保存結果，三台 HFSS 模擬機各自取件；人與 AI 判讀結果、決定下一步。",
    G4 / "arch.png", pad=0.05, trimmed=True,
    notes=[
        "這張是 10/12–10/23 寬頻濾波器任務實際在跑的流程；上一頁是時間軸，講誰在等誰，這頁講系統由哪些元件組成、資料怎麼流，顏色沿用上一頁的四種功能色。",
        "一輪：候選產生一次產生大量候選，代理模型或實測統計排序挑出約 24 筆，查重後分三份放進佇列；三台模擬機依優先順序各自取件跑 HFSS，輸出 S11、S21、S22 曲線，寫檔時用同一把尺算出最差餘裕。出錯會重試，同一筆錯三次就跳過。",
        "代理模型每輪用全部結果現場重訓；前代累積的約 2.1 萬筆濾波器資料拿來預訓。圖上把前代資料畫在共享儲存，實際放在開發機，屬簡化。",
        "好設計至少換一台機器重跑，一致才記為紀錄；交付候選四筆三台一致，曲線最大差不超過 2×10⁻⁶ dB。曾有一台對少數幾何算出不同結果，之後以另外兩台為準。",
        "規格由學長修訂；研究者決定目標與停止點，例如先做阻帶 −15 dB 版、到 +0.10 dB 才切回 −20 dB 版；AI 助手提出假設、寫程式、排下一輪並判讀。每一輪開始前先寫下成功與否的判準。",
        "這是研究用工具，不是量產系統；所有數字是 HFSS 模擬。",
    ])

# ================================================================ 27 第四代嘗試歷程：設計手法與代理模型
attempts_slide(
    "第四代嘗試歷程：設計手法與代理模型",
    "7–8 月在單埠天線上比較不同的候選來源：結構分析、演化、代理模型、局部修改與可微物理模擬，並複驗模擬設定。",
    [["7/06–7/15", "逐塊遮蔽，找出各塊金屬的作用", "兩片翼貢獻帶內約 6 dB", "留作分析工具"],
     ["7/11", "演化加代理模型粗篩", "候選逐代變窄，合格數逐代下降", "改成隨機生成加代理模型篩選"],
     ["7/13–7/29", "代理模型換新架構、修正訓練設定", "測試誤差約減半，仍只適合粗篩", "每輪保留隨機與模型分歧的候選"],
     ["7/25–7/31", "以金屬塊為單位修改、移植局部結構", "移植後小步修改到 +0.79 dB", "成為後段主要手法"],
     ["8/02–8/04", "可微物理模擬", "改善方向只猜對 60%，未達事先寫定的 70%", "不用於生成，暫停"],
     ["8/03–8/04", "較嚴的模擬網格複驗", "抽測 3 個通過設計全轉不通過", "送板前需要複驗"]],
    notes=[
        "結構分析：逐塊遮蔽金屬、拿掉翼、比較 3–6 塊的組數；兩片翼合計貢獻帶內約 6 dB，拿掉翼時方向圖反而變好，4–5 塊是較好的組數。這段一週內從 −2.68 推到 +0.35 dB，在不同機器重跑一致；之後把設計知識寫進候選空間的定義，不靠人挑。",
        "演化：7/11 演化加代理模型粗篩三代，合格數 12、10、6 逐代下降，候選逐代變窄，之後改成隨機生成加代理模型篩選。",
        "代理模型：7/23 資料約 1.45 萬筆加鏡射擴增後，新架構的凍結測試誤差 0.54，原架構 1.21，實戰連續勝出才換上；7/29 修正訓練的學習率設定，凍結測試誤差約減半。好設計之間仍排不準，所以每輪保留隨機與模型最不確定的候選。",
        "局部修改：先把另一款設計的局部結構移植到主流骨架，得 −0.15 dB，再連續小步修改到 +0.79 dB，在不同機器重跑三次一致。七巧板式的固定形狀庫第四代沒有做；相近的做法是以金屬塊為單位修改與區塊移植，仍在 25×25 像素上操作。",
        "可微物理模擬：自建可微的電磁模型，從腔模型做到分層介質 Green 函數，想當評分器並直接用梯度生成圖形。在資料沒覆蓋的形態上排序相關 +0.43，代理模型約 +0.07；事先寫定單像素翻轉的改善方向要猜對 70% 才用於生成，實際 60%，所以暫停，也沒有取代代理模型。",
        "較嚴的網格設定是收斂門檻 0.005、最多 20 次細化；抽測 3 個通過的天線設計，+0.45、+0.29、+0.13 dB 分別變成 −0.55、−0.95、−0.70 dB；+0.79 的設計變成 −0.25 dB。天線數字只在目前的模擬設定下成立。",
    ])

# ================================================================ 28 第四代嘗試歷程：對稱、饋入與新任務
attempts_slide(
    "第四代嘗試歷程：對稱、饋入與新任務",
    "8–10 月把同一套流程延伸到雙埠濾波器、對稱天線與寬頻濾波器，對稱與饋入金屬的做法也隨任務調整。",
    [["7 月", "前代散亂設計局部左右對稱化", "碎片消失，−1.42 → −0.29 dB", "之後首次通過規格"],
     ["8/06–8/11", "好設計事後完全鏡射", "天線中位下降 17.16 dB", "對稱要在生成時內建"],
     ["8/10–8/24", "雙埠濾波器：上下對稱生成、人工結構、無人值守迴圈", "人工結構不如演化出的設計", "未通過，9 月起暫停"],
     ["10/07–10/15", "對稱天線：生成時鏡射，再成對修改", "5,000 筆只 1 筆邊緣通過；成對修改後 +0.60 dB", "通過，目前的模擬設定"],
     ["10/12–10/23", "寬頻濾波器：兩埠金屬斷開", "11 GHz 穿透由約 −3 dB 變約 −18 dB", "只搜兩埠斷開的設計"],
     ["全程", "饋入固定金屬", "單埠在下緣中央；雙埠上下各 5×5 格", "四代都保留"]],
    notes=[
        "對稱 7 月：一個前代散亂設計做局部左右對稱化，−1.42 升到 −0.29 dB，碎片消失，在不同機器重跑一致；再從它局部修改，單埠天線首次通過規格。另把 8 個好設計全對稱化，方向圖 8 個都變好、帶內 7 個變差。",
        "對稱 8 月：13 個已通過規格的天線設計做三種鏡射，共 34 筆，每個設計模擬一次，中位下降 17.16 dB；濾波器事後對稱化中位下降 9.98 dB。濾波器的對照顯示生成時就對稱與不對稱的品質相當，所以濾波器把上下對稱生成設為預設，S11 與 S22 因此相同。",
        "雙埠濾波器：8/13 發現當時的模擬幾何把只碰一點的對角接點當成導通，改用與天線相同的可製造幾何。人工結構試過開縫、縫寬 0.03–0.15 mm 連續掃描、蛇形與摺疊諧振器。8/21 起無人值守迴圈：腳本自動挑選、派工、重跑驗證，迴圈裡沒有 AI 判斷，破紀錄只寫待審，由人確認後才算數；8/23 的 −2.43 dB 與 8/24 的 −2.39 dB 由它產生。",
        "對稱天線：饋入在下緣中央，天線只能左右對稱。+0.02 dB 那一筆重跑為 +0.04 dB，屬邊緣通過；成對修改是從好設計一次翻一組左右對稱的像素，+0.60 dB 含 0.10 mm 菱形橋，在不同機器重跑一致，方向圖不列入這個任務的評分。",
        "寬頻濾波器：八月的好設計改在 11–40 GHz 模擬，13.5–14 GHz 都有寄生通帶；兩埠相連時 11 GHz 幾乎直通，斷開才壓得下去，所以只搜兩埠斷開的設計。10/20 試的上下 180° 旋轉對稱 14 筆，最佳 −7.97 dB，撤掉。",
        "饋入：單埠的饋點永遠是金屬，修整後饋電那塊若不到 4 像素，就在下緣中央補一塊 3×3 金屬；雙埠上下緣中央各固定一塊 5×5 格金屬，生成後一律蓋回。生成器改成直接優化圖形時，與饋入相連的金屬比例由 0.62 掉到約 0.20，第四代沒有把連到饋入當成必要條件。",
    ])

# ================================================================ 29–30 新版代理模型
fig_slide(
    "新版代理模型：預測修改後的變化",
    "原本的代理模型從一張圖形預測整條曲線；新版拿量過的相近設計當參考，只預測改了哪幾格之後曲線怎麼變。",
    G4 / "surrogate_idea.png", pad=0.3,
    notes=[
        "示意圖，格子與曲線不是實際設計。",
        "原本：5 個多層感知器，輸入整張 25×25 圖形，輸出 S11、S21、S22 整條曲線，每輪重訓。變化量模型：卷積網路，輸入相近設計的圖形、改動的格子與它的 HFSS 曲線，輸出曲線變化量。上下文模型：同時看 1–4 個相近設計，用注意力機制互相參考，逐頻點決定信任哪個鄰居。",
        "動機：原本的模型學的是整條曲線，改幾格造成的變化只是兩次預測相減，誤差疊在一起；選樣當下預測與實測的排序相關約 0.10–0.40，平均偏樂觀 3.60–4.40 dB。",
        "訓練資料：前代約 2.1 萬筆加寬頻實測，配對距離取 1–50 格內，每個訓練回合重抽配對。新資料進來可以直接當參考，不必重訓就能用上。上下文模型訓練時混用 1–4 個鄰居，只給一個鄰居時排序也比專門的單鄰居模型好。",
        "離線回測：歷史留出資料上，改 12 格以內的排序相關由 0.40 升到 0.60；時間切分回測整體排序 0.36 升到 0.45–0.49、平均誤差 3.15 降到 2.30–2.39 dB。事先訂的門檻之一沒過，所以先小規模上線試。",
    ])

fig_text_slide(
    "新版代理模型：上線效果",
    "上線 10 輪，變化量模型的預測排序每輪都高於原本的代理模型；到最好的一群設計之間，排序能力消失。",
    G4 / "surrogate_rank.png",
    [("排序", ["變化量模型 0.24 到 0.71", "原本的代理模型 −0.63 到 +0.43，10 輪中 8 輪為負"]),
     ("最差餘裕高於 −5 dB 的比例", ["局部修改的候選：80%，原本 53%", "產生新設計的候選：57%，原本 19%"]),
     ("限制", ["最好的一群設計彼此差距小於模型解析度", "10/21 起改依實測統計挑局部修改"])],
    img_w=5.6, size=16,
    notes=[
        "縱軸是選樣當下、結果回來前的預測，與 HFSS 最差餘裕的 Spearman 排序相關。第 1–3 輪只算同一輪 10 筆從好設計局部修改的候選，兩個模型各挑 5 筆；第 4–10 輪算整輪 23–24 筆。第 8–10 輪改以阻帶 −15 dB 版計分，之前是 −20 dB 版。橫軸只列有逐輪紀錄的 10 輪。上下文模型第 7 輪起同步預測、先不參與挑選。",
        "80% 對 53%：同三輪、各 15 筆，阻帶 −20 dB 版；最佳值兩者相近，−4.07 對 −3.79 dB，最佳反而出在原本的模型挑的高變異設計。之後兩輪新版最佳與中位都不輸，局部修改的候選全交給新版。",
        "57% 對 19%：新版 4 輪 21 筆，對照是之前 9 輪原本模型挑的 53 筆，不是同一輪比較；新版這一塊出了當時全體最佳 −3.11 dB，只模擬一次。",
        "兩個新版模型比較：同一輪局部修改的候選，上下文模型單獨打分中位 −0.38、最佳 −0.25 dB；兩模型平均 −0.40、−0.29 dB，各 15 筆，差距小。",
        "實測高於 −1.50 dB 的 57 筆中，三種打分的排序相關只有 +0.05、+0.16、+0.08。10/21 起改用實測統計：把好設計上翻單一格的實測變化平均起來，假設可相加，預估組合的效果；加上好設計之間移植區塊，從 −0.20 推到 +0.41 dB，最後一段在多台機器重跑一致。實測統計的預估也偏樂觀，組合常不可相加。",
    ])

# ================================================================ 31–34 第四代：四任務爬升圖
fig_slide(
    "單埠天線：最差餘裕的爬升",
    "從前代累積的設計出發：27,433 筆模擬、29 天，從 −2.68 到 +0.79 dB；加上菱形橋的可製造版本為 +0.58 dB。",
    G4 / "climb_single.png",
    notes=[
        "縱軸是最差餘裕，只算 26.5–29.5 GHz 的 S11 與增益；藍點是方向圖與可製造條件都通過的設計，灰點是沒通過的，累計最佳線只計藍點；低於 −4 dB 的點畫在底線上。",
        "起點 −2.68 dB：前代資料中的一個設計拔掉 1–3 像素小碎片後的版本，7/03 第一輪模擬中的最佳值。",
        "第 346 筆首次通過，+0.15 dB：部分左右對稱化讓碎片消失，再從好設計局部修改。中段 +0.49 到 +0.56 dB：大量產生好設計的變體，代理模型初篩後模擬。終點 +0.79 dB，7/31：移植另一款設計的局部結構，再局部修改；方向圖餘裕 +0.20 dB，在不同機器重跑三次結果一致。",
        "可製造版本：同一設計加 0.10 mm 菱形橋為 +0.58 dB，在不同機器重跑一致。以更嚴的求解器網格設定模擬時，這個設計是 −0.25 dB，數字不能當成量測預測。",
        "筆數只算不重複圖形，含模擬機空閒時自動產生的隨機設計，主動挑選的是 16,680 筆；期間 7/03–7/31。全部是 HFSS 模擬。",
    ])

fig_slide(
    "對稱天線：最差餘裕的爬升",
    "從既有天線設計出發，金屬精確左右鏡射：5,144 筆模擬、9 天，從 −3.53 到 +0.60 dB。方向圖不列入這個任務的評分，照單埠的規格算是 +0.81 dB。",
    G4 / "climb_sym.png",
    notes=[
        "生成時金屬精確左右鏡射，可自由變動的像素由 624 減為 324，對角接點全部加 0.10 mm 菱形橋。",
        "起點 −3.53 dB：以我們的單埠好設計左右鏡射當種子，第一輪 16 筆中的最佳；第二輪由代理模型挑選後升到 −0.32 dB。",
        "約第 2,240 筆首次通過，約 +0.04 dB，屬邊緣通過，前 5,000 筆只有這一筆。終點 +0.60 dB，10/15：從好設計一次翻一組左右對稱的像素；在不同機器重跑結果一致，S11 餘裕 +0.70、增益餘裕 +0.60 dB。",
        "評分只看 S11 與正向增益，方向圖不列入這個任務的評分，照單埠的規格算是 +0.81 dB。10/12–10/14 沒有模擬對稱天線。10/15 之後另有只跑過一次、尚未重跑確認的較高值，不列入。全部是 HFSS 模擬。",
    ])

fig_slide(
    "雙埠濾波器：最差餘裕的爬升",
    "從前代累積的設計出發：15,625 筆模擬、12 天，從 −5.90 到 −2.39 dB，未通過規格；卡在通帶邊緣與阻帶起點，目前暫停。8/31 交付的 14 個濾波器也都未通過。",
    G4 / "climb_dual.png",
    notes=[
        "規格：24–32 GHz 每 0.5 GHz 一點；S11、S22 ≤ −10 dB，S21 通帶 25.5–30.5 GHz ≥ −3 dB，阻帶 24–25 與 31–32 GHz ≤ −15 dB；可製造幾何、對角接點加 0.075 mm 菱形橋。低於 −9 dB 的點畫在底線上。",
        "起點 −5.90 dB，8/13：前代雙埠設計原樣換成可製造幾何並加橋，第一輪 90 筆中的最佳。",
        "從好設計局部修改：同一天每步翻 1–2 個像素，到 −4.26、−3.82 dB。代理模型挑選、模擬機不停機補量：8/16 到 −2.88 dB，8/18 一天內刷新四次。自動化流程：8/21 起腳本自動挑選與重跑驗證，8/23 的 −2.43 dB 與 8/24 的 −2.39 dB 兩次刷新由它產生；終點在不同機器重跑一致。",
        "8/13 以前的舊模擬幾何把零厚度的對角點接觸當成導通，那段的數字不放在這張圖。0.075 mm 橋低於一般 PCB 最小線寬 0.1 mm，請日月光確認。",
        "筆數含空閒時的隨機設計，主動挑選 8,450 筆。全部是 HFSS 模擬。",
    ])

fig_slide(
    "寬頻濾波器：最差餘裕的爬升",
    "從既有濾波器設計出發，阻帶 −15 dB 版：2,252 筆模擬、9 天，從 −10.99 到 +0.41 dB。",
    G4 / "climb_wideband.png",
    notes=[
        "規格：11–40 GHz；S21 通帶 26–30 GHz ≥ −3 dB、阻帶 11–19 與 36–40 GHz ≤ −15 dB；S11、S22 在 26.5–29.5 GHz ≤ −10 dB；對角接點加 0.075 mm 菱形橋。",
        "前段以阻帶 −20 dB 版為挑選目標，圖上全部統一用 −15 dB 版計分。−20 dB 版目前最佳約 −3.05 dB，只跑過一次、尚未重跑確認，卡在 11–19 GHz 低阻帶。",
        "起點 −10.99 dB：前幾版規格下我們自己的好設計改用 11–40 GHz 重新模擬，第一輪 24 筆中的最佳；這些設計在 13.5–14 GHz 都有寄生通帶。",
        "兩埠金屬斷開、大步變異：一次改 30 格，把寄生峰壓下去，到 −5.75 dB。代理模型改學「改哪幾格、曲線怎麼變」，用來排序候選，到 −1.76 dB。好設計之間互相移植區塊，第 1,371 筆首次通過，約 +0.01 dB。通帶兩側成對修改，同時抬起 26 與 30 GHz 兩緣，到 +0.36 dB；再移植到其他好設計，終點 +0.41 dB，10/23，在不同機器重跑曲線完全一致。",
        "求解器設定與天線同等級，沒有做網格收斂檢查。寬頻濾波器每筆模擬中位約 464 秒。全部是 HFSS 模擬。",
    ])

# ================================================================ 35 量化總表
table_slide(
    "第四代：四個任務的量化總表",
    "所有數字為目前的模擬設定下的 HFSS 模擬，模擬網格尚未做收斂驗證；寬頻濾波器以阻帶 −15 dB 版計分。",
    [["任務", "起點最差餘裕", "終點最差餘裕", "模擬筆數", "天數", "狀態"],
     ["單埠天線", "−2.68 dB", "+0.79 dB\n加橋 +0.58 dB", "27,433", "29 天", "通過\n目前的模擬設定"],
     ["對稱天線", "−3.53 dB", "+0.60 dB", "5,144", "9 天", "通過\n目前的模擬設定"],
     ["雙埠濾波器", "−5.90 dB", "−2.39 dB", "15,625", "12 天", "未通過"],
     ["寬頻濾波器", "−10.99 dB", "+0.41 dB", "2,252", "9 天", "通過\n目前的模擬設定"]],
    [2.45, 2.0, 2.3, 1.85, 1.4, 2.17], row_h=0.92, header_size=16, body_size=18,
    center_cols=(1, 2, 3, 4, 5), header_h=0.55,
    notes=[
        "口徑：最差餘裕是每條規格在每個評分頻點離門檻多少 dB、取最差一點，正值代表全部通過；模擬筆數只算不重複的圖形；起點是第一輪模擬中的最佳值；天數是第一筆到終點那一筆的日曆天數，含頭尾。",
        "期間：單埠天線 7/03–7/31；對稱天線 10/07–10/15；雙埠濾波器 8/13–8/24；寬頻濾波器 10/15–10/23。四個終點都在不同機器重跑結果一致。",
        "起點來源：單埠天線與雙埠濾波器取自前代資料；對稱天線與寬頻濾波器是我們較早的設計。模擬筆數含模擬機空閒時自動產生的隨機設計，主動挑選的是單埠 16,680、雙埠 8,450。",
        "條件：單埠天線終點無菱形橋，加 0.10 mm 橋後 +0.58 dB；對稱天線含 0.10 mm 橋；兩個濾波器含 0.075 mm 橋。以更嚴的求解器網格設定模擬時，單埠天線 +0.79 的設計變成 −0.25 dB，抽測的 3 個通過設計都轉為不通過，送板前需要複驗。",
        "寬頻濾波器阻帶 −20 dB 版尚未通過。",
    ])

# ================================================================ 36–38 最佳設計
fig_slide(
    "最佳設計：單埠天線",
    "最差餘裕 +0.79 dB、方向圖餘裕 +0.20 dB；加上 0.10 mm 菱形橋的可製造版本為 +0.58 dB。HFSS 模擬，未實測。",
    G4 / "best_single.png",
    notes=[
        "由左至右：金屬圖形，黑色是金屬、饋點在下緣中央；S11 與正向增益，虛線是門檻、淺紅是評分頻段；28 GHz 方向圖，相對 0° 方向，淺紅是 ±45° 內低於 0° 方向 3 dB 的禁區、虛線是 −3 dB 門檻。",
        "2026/07/31 的設計，在不同機器重跑三次結果一致；圖為無橋幾何。以更嚴的求解器網格設定模擬時為 −0.25 dB。",
        "8/31 交付的天線是同款式中加橋後的代表設計，不是這個設計本身。",
    ])

fig_slide(
    "最佳設計：對稱天線",
    "金屬精確左右鏡射，對角接點加 0.10 mm 菱形橋，最差餘裕 +0.60 dB；方向圖不列入這個任務的評分，照單埠的規格算是 +0.81 dB。HFSS 模擬，未實測。",
    G4 / "best_sym.png",
    notes=[
        "由左至右：金屬圖形，黑色是金屬、饋點在下緣中央、左右鏡射；S11 與正向增益，S11 餘裕 +0.70、增益餘裕 +0.60 dB；28 GHz 方向圖，相對 0° 方向。",
        "方向圖 +0.81 dB 是兩切面中較差的 φ＝0°，φ＝90° 為 +1.29 dB，用單埠天線同一套 ±45°、−3 dB 的規格現場計算。",
        "2026/10/15 的設計，在兩台不同機器重跑，響應完全一致；這只代表數值可重現，不代表與實測相符。",
        "這個設計不在 8/31 交付內；是否加入下次送板，請日月光決定。",
    ])

fig_slide(
    "最佳設計：寬頻濾波器",
    "阻帶 −15 dB 版通過，最差餘裕 +0.41 dB，最差的一項是通帶；阻帶 −20 dB 版尚未通過。兩埠金屬不相連，HFSS 模擬，未實測。",
    G4 / "best_wideband.png",
    notes=[
        "由左至右：金屬圖形，上下緣中央是兩埠的接點；S21、S11、S22，虛線是阻帶 −15 dB 版的門檻、淺紅是評分頻段。",
        "各項餘裕：通帶 +0.41、低阻帶 +1.04、高阻帶 +2.95、S11 +1.24、S22 +4.71 dB；通帶 26–30 GHz 內 S21 最差一點約 −2.59 dB。",
        "同一設計以阻帶 −20 dB 版重新評分是 −3.96 dB，卡在 11–19 GHz 低阻帶。2026/10/23 的設計，在兩台不同機器重跑，曲線差為 0；另有次選設計 +0.40 dB，也在不同機器重跑一致。",
        "兩埠金屬斷開的理由：11 GHz 時兩埠金屬相連的設計 S21 約 −3 dB，幾乎直通；斷開的約 −18 dB。對角接點加 0.075 mm 菱形橋，低於一般 PCB 最小線寬 0.1 mm，請日月光確認。",
    ])

# ================================================================ 39 可製造：第四代
fig_table_slide(
    "第四代的可製造：懸浮金屬、對角導通與菱形橋",
    "懸浮金屬與小碎片參與共振，第四代不刪也不接；出圖時在對角接點加菱形橋，天線 0.10 mm、濾波器 0.075 mm。",
    MEET / "schem_diamond.png",
    [["做法", "結果"],
     ["拔掉 1–3 像素的小碎片", "5 個前代設計中 4 個下降 4.67–16.93 dB"],
     ["懸浮金屬用 1 像素接回主體", "6 個試驗中 4 個大幅下降"],
     ["事後清除對角接點", "三種做法響應都大幅變差"],
     ["對角接點加 0.10 mm 菱形橋", "3,319 個天線設計，中位變化 −0.07 dB"],
     ["對角接點直接切斷", "同樣這些設計，中位下降 4.84 dB"],
     [("請日月光確認", True), "低於一般 PCB 0.1 mm 的細橋製程規則；8/31 交付中有 18 個設計的橋寬低於 0.1 mm"]],
    [2.7, 4.15], img_w=5.0, row_h=0.66, body_size=14,
    notes=[
        "左圖照實際模擬幾何繪製：原始幾何的對角接點只靠 0.01 mm 重疊相連；直接切斷是在角點挖掉一顆 0.10 mm 的 45° 菱形；加橋是在角點補一顆同尺寸的菱形金屬。示意圖不代表特定設計。",
        "小碎片：把前代資料中 5 個通過或接近通過的設計拔掉 1–3 像素小碎片，4 個下降 4.67–16.93 dB；只有整塊型的一個幾乎不變，成為單埠天線的可製造起點 −2.68 dB。懸浮金屬：用 1 像素把懸浮金屬接回主體，6 個試驗中有 4 個的最差餘裕掉到 −5 至 −6 dB 之間。小碎片與懸浮金屬參與共振，不能事後直接刪或接；可製造改定義為每塊金屬至少 4 像素。",
        "菱形橋：2026/08/05–08/10 對 3,319 個天線設計各模擬加橋與切斷兩種幾何，每個設計各模擬一次；加橋後 86.50% 的設計變化在 0.30 dB 以內，對角接點 17 座以上時這個比例降到約四成。單埠天線 +0.79 的設計加橋後為 +0.58 dB，在不同機器重跑一致。",
        "濾波器的舊模擬幾何把零厚度的對角點接觸當成導通，8/13 起改用與天線相同的可製造幾何；同一個濾波器設計切掉橋、0.01 mm 微橋、0.075 mm 橋分別是 −31.03、−6.86、−3.82 dB。",
        "前三代的做法分在各代段落：第一代口試後提出四邊各擴 0.05 mm 的構想；第二代用 0.21 mm 重疊填入並保留孤島；第三代在訓練中用連通 loss。",
    ])

# ================================================================ 收尾：40 現況與下一步
table_slide(
    "現況與下一步",
    "截至 2026/10/23，所有第四代數字都是 HFSS 模擬；右兩欄是可討論的選項與請日月光決定的事，尚未決定。",
    [["項目", "現況", "可討論的選項", "請日月光決定"],
     ["單埠天線", "通過，+0.79 dB；加 0.10 mm 橋 +0.58 dB", "打樣量測", "—"],
     ["對稱天線", "通過，+0.60 dB，含 0.10 mm 橋", "打樣量測", "是否加入下次送板"],
     ["雙埠濾波器", "未通過，−2.39 dB，目前暫停", "依需求再啟動", "—"],
     ["寬頻濾波器", "阻帶 −15 dB 版通過，+0.41 dB；−20 dB 版未通過，最佳約 −3.05 dB", "送板打樣，或繼續推進低阻帶",
      "阻帶要 −15 還是 −20 dB"],
     ["8/31 交付", "39 個設計：天線 25、濾波器 14，尚未實測；濾波器全部未通過規格", "打樣量測", "細橋製程規則"],
     ["模擬與實測", "較嚴的模擬設定下，抽測 3 個通過設計全轉不通過；第三代實測通帶下緣比模擬高約 0.5–0.8 GHz",
      "送板前以較嚴設定複驗；量測後回饋設計", "—"]],
    [1.75, 4.75, 3.0, 2.67], row_h=0.66, body_size=14,
    notes=[
        "模擬限制：所有第四代數字使用同一套 HFSS 求解器設定，尚未做網格收斂。以更嚴的網格設定模擬時共振頻率會移動，單埠天線 +0.79 的設計變成 −0.25 dB，抽測的 3 個通過設計都轉為不通過。",
        "實測偏移：第三代 5 個樣品的量測下緣比模擬高 0.5–0.8 GHz、上緣高 1.4–1.6 GHz，出自 2026/03/24 月報。第四代的設計還沒有實測，模擬與量測會差多少目前沒有數據，也還沒把這個偏移納入設計。",
        "8/31 交付：天線 25 個，含 7 個款式的代表 21 個與前代資料中的頂尖設計 4 個；濾波器 14 個主要卡在 S21 通帶邊緣與阻帶起點。39 個設計全部在不同機器重跑，37 個完全一致、2 個差 0.08 dB 以內；18 個設計的橋寬低於一般 PCB 最小線寬 0.1 mm。",
        "寬頻濾波器 −3.05 dB 只跑過一次、尚未重跑確認；若維持 −20 dB，需要新的兩埠斷開款式當起點。對稱天線 10/15 之後另有只跑過一次的較高值，不列入。",
        "對稱天線與寬頻濾波器都是 8/31 之後的結果，不在已交付的 39 個設計內。以上都是選項，供討論，不代表已決定。",
    ])

# ================================================================ 41 附錄：名詞白話表
s, top = content_slide("附錄：名詞白話表", "第一次讀到不懂的詞先查這裡；左欄是評分與設計，中欄是前三代的方法，右欄是第四代的方法。", notes=[
    "名詞以本報告的用法為準；各代方法名稱沿用碩論原名。"])
TERMS = [
    [("像素化金屬圖形", "25×25 格逐格填金屬或留空"),
     ("HFSS", "電磁全波模擬軟體，四代共用"),
     ("最差餘裕", "各規格各頻點離門檻，取最差"),
     ("方向圖", "28 GHz 兩切面的輻射強度"),
     ("0° 方向", "天線面法線方向，方向圖基準"),
     ("目前的模擬設定", "第四代共用的求解器設定"),
     ("饋入固定金屬", "饋入點附近強制放的金屬"),
     ("孤島、懸浮金屬", "沒有連到饋入點的金屬"),
     ("饋入連通率", "與饋入點相連的金屬占比"),
     ("0.21 mm 重疊填入", "每格畫大一點，讓對角重疊"),
     ("菱形橋", "對角接點補一顆 45° 菱形"),
     ("可製造", "每塊至少 4 像素、對角加橋"),
     ("寄生通帶", "阻帶裡意外出現的穿透峰")],
    [("代理模型", "近似 HFSS 的神經網路"),
     ("線上學習", "生成、模擬、訓練交替進行"),
     ("強化學習", "依獎勵逐步調整動作"),
     ("上下界 loss", "只罰超出上下界的部分"),
     ("BiScaleNorm", "正負兩側分開縮放"),
     ("二階段生成", "先收斂，再週期翻轉像素"),
     ("回滾", "沒改善就退回歷史最佳圖"),
     ("小圖案拼接", "每格從幾種小圖案挑一個"),
     ("七巧板", "從形狀庫挑小形狀拼滿"),
     ("二值化溫度", "控制輸出接近 0 或 1 的程度"),
     ("ACP", "停滯時拉高學習率與溫度"),
     ("DLF", "只複習優於累計平均的樣本"),
     ("孤島抑制 loss", "SC、TV、IS、GC 四種")],
    [("批次線上學習", "一組候選送三台，收齊再挑"),
     ("局部修改", "從好設計只改少數像素"),
     ("區塊移植", "好設計的一塊金屬換到別處"),
     ("成對修改", "一次改一組，例如左右一對"),
     ("多台機器重跑一致", "換一台機器重跑，響應相同"),
     ("無人值守迴圈", "腳本自動挑選、派工、重跑"),
     ("變化量模型", "預測改格後的曲線變化"),
     ("上下文模型", "同時參考 1–4 個相近設計"),
     ("實測統計預估", "用單格實測變化估組合"),
     ("排序相關", "預測與實測排名的一致程度"),
     ("可微物理模擬", "自建可微的電磁模型"),
     ("兩埠金屬斷開", "濾波器兩埠金屬互不相連"),
     ("雙埠濾波器", "上下各一埠，看穿透與反射")],
]
gap = 0.22
tw = (CW - 2 * gap) / 3
for k, terms in enumerate(TERMS):
    table(s, L + k * (tw + gap), top + 0.02, [["名詞", "白話"]] + [[a, b] for a, b in terms], [1.5, tw - 1.5],
          header_size=12, body_size=11, row_h=0.35, header_h=0.36)

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
